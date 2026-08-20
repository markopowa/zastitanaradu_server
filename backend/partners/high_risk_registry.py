import io

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from .models import ClientCompany, Employee


def _write(cell, text, bold=False, size=9, center=False):
    cell.text = ""
    p = cell.paragraphs[0]
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(str(text) if text is not None else "")
    r.bold = bold
    r.font.size = Pt(size)


_HEADER = [
    "Redni\nbroj",
    "Naziv radnog mesta sa povećanim rizikom",
    "Ime i prezime zaposlenog",
    "Nivo rizika",
]

_COL_WIDTHS = [1.5, 8.0, 6.0, 4.0]


def generate_high_risk_registry(client_id: int) -> bytes:
    company = ClientCompany.objects.get(pk=client_id)

    employees = (
        Employee.objects.filter(client_company_id=client_id)
        .select_related("job_role", "job_role__risk_level", "risk_level_override")
        .order_by("last_name", "first_name")
    )

    rows = [
        emp
        for emp in employees
        if emp.effective_risk_level is not None
        and emp.effective_risk_level.is_high_risk
    ]

    doc = docx.Document()
    section = doc.sections[0]
    section.left_margin = section.right_margin = Cm(1.5)
    section.top_margin = section.bottom_margin = Cm(1.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t = title.add_run("EVIDENCIJA O RADNIM MESTIMA SA POVEĆANIM RIZIKOM")
    t.bold = True
    t.font.size = Pt(12)

    doc.add_paragraph()

    info = doc.add_table(rows=2, cols=3)
    info.style = "Table Grid"
    _write(info.cell(0, 0), company.name, bold=True)
    _write(info.cell(0, 1), company.address or "")
    _write(info.cell(0, 2), company.tax_id or "")
    _write(info.cell(1, 0), "Poslovno ime poslodavca", size=8)
    _write(info.cell(1, 1), "Adresa sedišta poslodavca", size=8)
    _write(info.cell(1, 2), "PIB poslodavca", size=8)

    doc.add_paragraph()

    tbl = doc.add_table(rows=1 + len(rows), cols=4)
    tbl.style = "Table Grid"

    for col_idx, width in enumerate(_COL_WIDTHS):
        for cell in tbl.columns[col_idx].cells:
            cell.width = Cm(width)

    for i, text in enumerate(_HEADER):
        _write(tbl.cell(0, i), text, bold=True, size=9, center=True)

    for ordinal, emp in enumerate(rows, start=1):
        row = tbl.rows[ordinal]
        _write(row.cells[0], str(ordinal), center=True)
        _write(row.cells[1], emp.high_risk_position_name or (emp.job_role.name if emp.job_role_id else "") or "")
        _write(row.cells[2], f"{emp.first_name} {emp.last_name}".strip())
        _write(row.cells[3], emp.effective_risk_level.label)

    doc.add_paragraph()

    sig = doc.add_table(rows=1, cols=2)
    sig.cell(0, 0).text = (
        "_____________________________\n"
        "Savetnik/saradnik za bezbednost i zdravlje na radu"
    )
    sig.cell(0, 1).text = "_____________________________\nPoslodavac"

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()
