import json
from pathlib import Path

import docx

from django.conf import settings
from django.core.management.base import BaseCommand
from django.db import transaction

from documents.models import DocumentTemplate
from documents.tag_placement_migration import (
    derive_placements,
    has_tags,
    strip_tags_from_docx,
)
from documents.utils import invalidate_page_images
from partners.management.commands.add_setup import (
    MASTER_TEMPLATE_DIR_NAME,
    MASTER_TEMPLATE_FILES,
    MASTER_TEMPLATE_GENERATION_CONFIG,
)
from partners.models import JobRole, TrainingType

BLANK_ROLE_FIELDS = {
    "obrazac6": ("obrazac6_template", "obrazac6_fields"),
    "lzo_revers": ("lzo_revers_template", "lzo_revers_fields"),
}


def _series_source_map(generation_config: dict) -> dict:
    return {
        entry.get("table"): entry.get("source")
        for entry in (generation_config or {}).get("series") or []
        if entry.get("source")
    }


def _load_sidecar(path: Path) -> dict:
    if not path.exists():
        return {}
    try:
        return json.loads(path.read_text())
    except (ValueError, OSError):
        return {}


def _write_sidecar(path: Path, data: dict) -> None:
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2, sort_keys=True))


class Command(BaseCommand):
    help = (
        "One-time migration: derive VISUAL placement coordinates from the "
        "existing {{ tag }} positions (by rendering the badge preview and "
        "locating each label with PyMuPDF), then strip the tags from the "
        "source .docx files. Idempotent: files without tags are skipped."
    )

    def add_arguments(self, parser):
        parser.add_argument(
            "--files-only",
            action="store_true",
            help=(
                "Process the repo's master_templates/ and pilot_blanks/ "
                "files directly (no DB required) and write placements.json "
                "sidecars instead of writing to DocumentTemplate/JobRole/"
                "TrainingType rows."
            ),
        )
        parser.add_argument(
            "--dry-run",
            action="store_true",
            help="Report what would change without writing anything.",
        )

    def handle(self, *args, **options):
        dry_run = options["dry_run"]
        if options["files_only"]:
            self._run_files_only(dry_run)
        else:
            self._run_db(dry_run)

    # ------------------------------------------------------------------
    # DB mode
    # ------------------------------------------------------------------

    def _run_db(self, dry_run: bool):
        with transaction.atomic():
            for tpl in DocumentTemplate.objects.exclude(
                template_file=""
            ).exclude(template_file__isnull=True):
                self._migrate_master(tpl, dry_run)

            for role in JobRole.objects.all():
                for _kind, (file_field, fields_field) in BLANK_ROLE_FIELDS.items():
                    self._migrate_blank(
                        owner=role,
                        file_field_name=file_field,
                        fields_field_name=fields_field,
                        label=f"JobRole#{role.id} {file_field}",
                        dry_run=dry_run,
                    )

            for training_type in TrainingType.objects.all():
                self._migrate_blank(
                    owner=training_type,
                    file_field_name="potvrda_template",
                    fields_field_name="potvrda_fields",
                    label=f"TrainingType#{training_type.id} potvrda_template",
                    dry_run=dry_run,
                )

            if dry_run:
                transaction.set_rollback(True)

    def _migrate_master(self, tpl: DocumentTemplate, dry_run: bool) -> None:
        file_field = tpl.template_file
        name = getattr(file_field, "name", "") or ""
        if not name.lower().endswith(".docx"):
            return
        path = Path(file_field.path)
        if not path.exists():
            return

        document = docx.Document(str(path))
        if not has_tags(document):
            self.stdout.write(f"[skip] DocumentTemplate#{tpl.id} '{tpl.name}': no tags.")
            return

        series_map = _series_source_map(tpl.generation_config or {})
        placements, warnings = derive_placements(path, series_source_map=series_map)
        self._report(f"DocumentTemplate#{tpl.id} '{tpl.name}'", placements, warnings)
        if dry_run:
            return

        strip_tags_from_docx(path)
        tpl.generation_config = {"mode": "VISUAL", "placeholders": placements}
        tpl.save(update_fields=["generation_config"])
        invalidate_page_images(tpl.pk)

    def _migrate_blank(
        self, owner, file_field_name: str, fields_field_name: str, label: str, dry_run: bool,
    ) -> None:
        file_field = getattr(owner, file_field_name, None)
        if not file_field:
            return
        name = getattr(file_field, "name", "") or ""
        if not name.lower().endswith(".docx"):
            return
        path = Path(file_field.path)
        if not path.exists():
            return

        document = docx.Document(str(path))
        if not has_tags(document):
            self.stdout.write(f"[skip] {label}: no tags.")
            return

        placements, warnings = derive_placements(path)
        self._report(label, placements, warnings)
        if dry_run:
            return

        strip_tags_from_docx(path)
        setattr(owner, fields_field_name, placements)
        owner.save(update_fields=[fields_field_name])

    # ------------------------------------------------------------------
    # Files-only mode (no DB dependency for the templates being migrated)
    # ------------------------------------------------------------------

    def _run_files_only(self, dry_run: bool):
        master_dir = Path(settings.BASE_DIR) / MASTER_TEMPLATE_DIR_NAME
        self._migrate_master_files(master_dir, dry_run)

        blanks_root = Path(settings.BASE_DIR) / "pilot_blanks"
        for kind in BLANK_ROLE_FIELDS:
            self._migrate_blank_files(blanks_root / kind, dry_run)

    def _migrate_master_files(self, master_dir: Path, dry_run: bool) -> None:
        if not master_dir.is_dir():
            return
        sidecar_path = master_dir / "placements.json"
        sidecar = _load_sidecar(sidecar_path)
        filename_to_name = {v: k for k, v in MASTER_TEMPLATE_FILES.items()}
        changed = False

        for path in sorted(master_dir.glob("*.docx")):
            key = filename_to_name.get(path.name, path.stem)
            document = docx.Document(str(path))
            if not has_tags(document):
                self.stdout.write(f"[skip] master '{path.name}': no tags.")
                continue

            doc_name = filename_to_name.get(path.name)
            generation_config = MASTER_TEMPLATE_GENERATION_CONFIG.get(doc_name, {})
            series_map = _series_source_map(generation_config)
            placements, warnings = derive_placements(path, series_source_map=series_map)
            self._report(f"master file '{path.name}'", placements, warnings)
            if dry_run:
                continue

            strip_tags_from_docx(path)
            sidecar[key] = {"mode": "VISUAL", "placeholders": placements}
            changed = True

        if changed and not dry_run:
            _write_sidecar(sidecar_path, sidecar)
            self.stdout.write(self.style.SUCCESS(
                f"Wrote {sidecar_path}"))

    def _migrate_blank_files(self, blanks_dir: Path, dry_run: bool) -> None:
        if not blanks_dir.is_dir():
            return
        sidecar_path = blanks_dir / "placements.json"
        sidecar = _load_sidecar(sidecar_path)
        changed = False

        for path in sorted(blanks_dir.glob("*.docx")):
            document = docx.Document(str(path))
            if not has_tags(document):
                self.stdout.write(f"[skip] blank '{path.name}': no tags.")
                continue

            placements, warnings = derive_placements(path)
            self._report(f"blank '{path.name}'", placements, warnings)
            if dry_run:
                continue

            strip_tags_from_docx(path)
            sidecar[path.name] = placements
            changed = True

        if changed and not dry_run:
            _write_sidecar(sidecar_path, sidecar)
            self.stdout.write(self.style.SUCCESS(
                f"Wrote {sidecar_path}"))

    # ------------------------------------------------------------------

    def _report(self, label: str, placements: list, warnings: list[str]) -> None:
        field_count = sum(1 for p in placements if "seriesSource" not in p)
        series_count = sum(1 for p in placements if "seriesSource" in p)
        self.stdout.write(
            f"[ok] {label}: {field_count} field(s), {series_count} series "
            f"group(s) derived, {len(warnings)} warning(s)."
        )
        for warning in warnings:
            self.stdout.write(self.style.WARNING(f"    {warning}"))
