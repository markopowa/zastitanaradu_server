import io
import uuid

import docx
from django.contrib.auth import get_user_model
from django.test import TestCase
from rest_framework.test import APITestCase

from partners.kinney import categorize, compute_rizik, is_high_risk
from partners.models import (
    ClientCompany,
    Hazard,
    JobRole,
    JobRoleHazard,
    RiskLevel,
)

User = get_user_model()


class KinneyUnitTest(TestCase):
    def test_compute_rizik_multiplies(self):
        self.assertEqual(compute_rizik(10, 6, 7), 420.0)
        self.assertEqual(compute_rizik(0.5, 1, 1), 0.5)

    def test_categorize_bands(self):
        self.assertEqual(categorize(10)[0], "PRIHVATLJIV")
        self.assertEqual(categorize(50)[0], "MOGUCI")
        self.assertEqual(categorize(150)[0], "ZNACAJAN")
        self.assertEqual(categorize(300)[0], "VISOK")
        self.assertEqual(categorize(900)[0], "VRLO_VISOK")

    def test_high_risk_flag(self):
        self.assertTrue(is_high_risk("VISOK"))
        self.assertTrue(is_high_risk("VRLO_VISOK"))
        self.assertFalse(is_high_risk("MOGUCI"))


class JobRoleHazardModelTest(TestCase):
    def setUp(self):
        self.company = ClientCompany.objects.create(
            name="Firma", tax_id=uuid.uuid4().hex[:9])
        self.role = JobRole.objects.create(
            client_company=self.company, name="Pekar")
        self.hazard = Hazard.objects.create(
            code="BUKA", label="Buka", kind=Hazard.KIND_STETNOST)

    def test_save_computes_rizik_and_category(self):
        jrh = JobRoleHazard.objects.create(
            job_role=self.role, hazard=self.hazard,
            verovatnoca=10, izlozenost=6, posledica=7)
        self.assertEqual(jrh.rizik, 420.0)
        self.assertEqual(jrh.risk_category, "VRLO_VISOK")

    def test_low_values_acceptable(self):
        jrh = JobRoleHazard.objects.create(
            job_role=self.role, hazard=self.hazard,
            verovatnoca=1, izlozenost=1, posledica=3)
        self.assertEqual(jrh.rizik, 3.0)
        self.assertEqual(jrh.risk_category, "PRIHVATLJIV")


def load_doc(content):
    return docx.Document(io.BytesIO(content))


class RiskAssessmentApiTest(APITestCase):
    def setUp(self):
        self.user = User.objects.create_superuser(
            "ra", "ra@example.com", "pass12345")
        self.client.force_authenticate(user=self.user)
        self.company = ClientCompany.objects.create(
            name="Pekara Klas", tax_id=uuid.uuid4().hex[:9])
        self.role = JobRole.objects.create(
            client_company=self.company, name="Pekar")
        self.hazard = Hazard.objects.create(
            code="PAD_VISINA", label="Pad sa visine",
            kind=Hazard.KIND_OPASNOST)

    def test_hazards_not_paginated_returns_all(self):
        for n in range(15):
            Hazard.objects.create(
                code=f"HZ{n}", label=f"Opasnost {n}",
                kind=Hazard.KIND_OPASNOST)
        r = self.client.get("/api/partners/hazards/")
        self.assertEqual(r.status_code, 200)
        self.assertIsInstance(r.data, list)
        self.assertGreaterEqual(len(r.data), 16)

    def test_kinney_scale_options_not_paginated(self):
        from partners.models import KinneyScaleOption
        for factor in ("V", "I", "P"):
            for i in range(6):
                KinneyScaleOption.objects.create(
                    factor=factor, value=float(i + 1) + 0.5,
                    label=f"{factor}{i}", order=i)
        r = self.client.get("/api/partners/kinney-scale-options/?factor=V")
        self.assertEqual(r.status_code, 200)
        self.assertIsInstance(r.data, list)
        self.assertEqual(len(r.data), 6)

    def test_create_hazard(self):
        r = self.client.post(
            "/api/partners/hazards/",
            {"code": "EL", "label": "Struja", "kind": "OPASNOST"},
            format="json")
        self.assertEqual(r.status_code, 201, r.data)

    def test_create_job_role_hazard_computes_rizik(self):
        r = self.client.post(
            "/api/partners/job-role-hazards/",
            {"job_role": self.role.id, "hazard": self.hazard.id,
             "verovatnoca": 6, "izlozenost": 6, "posledica": 15,
             "mere": "Zaštitna ograda"},
            format="json")
        self.assertEqual(r.status_code, 201, r.data)
        self.assertEqual(r.data["rizik"], 540.0)
        self.assertEqual(r.data["risk_category"], "VRLO_VISOK")

    def test_list_filtered_by_job_role(self):
        JobRoleHazard.objects.create(
            job_role=self.role, hazard=self.hazard,
            verovatnoca=1, izlozenost=1, posledica=1)
        r = self.client.get(
            f"/api/partners/job-role-hazards/?job_role_id={self.role.id}")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(len(r.data), 1)

    def test_generate_akt_contains_role_and_hazard(self):
        JobRoleHazard.objects.create(
            job_role=self.role, hazard=self.hazard,
            verovatnoca=6, izlozenost=6, posledica=15, mere="Ograda")
        r = self.client.get(
            f"/api/partners/client-companies/{self.company.id}"
            "/risk-assessment-act-generated/")
        self.assertEqual(r.status_code, 200)
        doc = load_doc(r.content)
        text = "\n".join(p.text for p in doc.paragraphs)
        table_text = "\n".join(
            cell.text for t in doc.tables for row in t.rows
            for cell in row.cells)
        self.assertIn("AKT O PROCENI RIZIKA", text)
        self.assertIn("Pekar", text)
        self.assertIn("Pad sa visine", table_text)
        self.assertIn("Pekar", "\n".join(
            [text]))
