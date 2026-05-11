import io
from collections import defaultdict

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from .models import ClientCompany
from processes.models import ProcessRun


def _fmt_date(d):
    if not d:
        return ""
    try:
        return d.strftime("%d.%m.%Y.")
    except AttributeError:
        return str(d)


def _result_data_field(rd, new_key, old_key):
    if not isinstance(rd, dict):
        return ""
    v = rd.get(new_key)
    if v is not None and str(v).strip() != "":
        return str(v)
    v2 = rd.get(old_key)
    return "" if v2 is None else str(v2)


def _write(cell, text, bold=False, size=9, center=False):
    cell.text = ""
    p = cell.paragraphs[0]
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(str(text) if text is not None else "")
    r.bold = bold
    r.font.size = Pt(size)


def _short_label(process_type_name: str) -> str:
    """First word of the process type name — e.g. 'Prethodni lekarski pregled' -> 'Prethodni'."""
    return process_type_name.split()[0] if process_type_name else ""


def _employee_rows(emp_runs: list) -> list:
    """
    Returns list of (process_type, run) tuples for one employee.
    Grouped by process type (sorted by name), runs within each type sorted by performed_at.
    """
    by_type: dict = defaultdict(list)
    for run in emp_runs:
        by_type[run.process_type].append(run)

    sorted_types = sorted(by_type.keys(), key=lambda pt: pt.name)
    rows = []
    for pt in sorted_types:
        for run in sorted(by_type[pt], key=lambda r: r.performed_at or ""):
            rows.append((pt, run))
    return rows if rows else [(None, None)]


_COL_WIDTHS = [1.2, 3.8, 3.2, 1.4, 1.8, 3.0, 3.0, 2.8, 2.5, 2.5]

_HEADER = [
    "Redni\nbroj",
    "Naziv radnog mesta sa povećanim rizikom koje je utvrđeno aktom o proceni rizika",
    "Ime i prezime zaposlenog koji radi na radnom mestu sa povećanim rizikom",
    "Interval vršenja periodičnih lekarskih pregleda izraženim u mesecima",
    "Datum izvršenog prethodnog i periodičnog lekarskog pregleda zaposlenog",
    None,
    "Datum kada treba da se izvrši sledeći lekarski pregled zaposlenog",
    "Broj lekarskog izveštaja",
    "Ocena zdravstvene sposobnosti",
    "Preduzete mere\n(raspoređen na drugo radno mesto - poslove)",
]


def generate_medical_exam_record(client_id: int) -> bytes:
    company = ClientCompany.objects.get(pk=client_id)

    runs = (
        ProcessRun.objects.filter(
            status=ProcessRun.STATUS_COMPLETED,
            process_binding__employee__client_company_id=client_id,
            process_type__include_in_medical_exam_record=True,
        )
        .select_related(
            "process_type",
            "process_binding",
            "process_binding__employee",
        )
        .order_by(
            "process_binding__employee__last_name",
            "process_binding__employee__first_name",
            "process_type__name",
            "performed_at",
        )
    )

    by_employee: dict = defaultdict(list)
    for run in runs:
        emp = run.process_binding.employee
        by_employee[emp].append(run)

    employee_data = [
        (emp, _employee_rows(emp_runs))
        for emp, emp_runs in by_employee.items()
    ]

    doc = docx.Document()
    section = doc.sections[0]
    section.page_width, section.page_height = section.page_height, section.page_width
    section.left_margin = section.right_margin = Cm(1.0)
    section.top_margin = section.bottom_margin = Cm(1.5)

    obr = doc.add_paragraph("OBRAZAC 1")
    obr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    obr.runs[0].font.size = Pt(9)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t = title.add_run(
        "EVIDENCIJA O RADNIM MESTIMA SA POVEĆANIM RIZIKOM, ZAPOSLENIMA KOJI OBAVLJAJU "
        "POSLOVE NA RADNIM MESTIMA SA POVEĆANIM RIZIKOM I LEKARSKIM PREGLEDIMA "
        "ZAPOSLENIH KOJI OBAVLJAJU TE POSLOVE"
    )
    t.bold = True
    t.font.size = Pt(10)

    doc.add_paragraph()

    info = doc.add_table(rows=2, cols=3)
    info.style = "Table Grid"
    _write(info.cell(0, 0), company.name, bold=True)
    _write(info.cell(0, 1), company.address or "")
    _write(info.cell(0, 2), company.tax_id or "")
    _write(info.cell(1, 0), "Poslovno ime poslodavca", size=8)
    _write(info.cell(1, 1), "Adresa sedišta poslodavca", size=8)
    _write(info.cell(1, 2), "PIB poslodavca", size=8)

    doc.add_paragraph()

    total_rows = 1 + sum(len(rows) for _, rows in employee_data)

    tbl = doc.add_table(rows=total_rows, cols=10)
    tbl.style = "Table Grid"

    for col_idx, width in enumerate(_COL_WIDTHS):
        for cell in tbl.columns[col_idx].cells:
            cell.width = Cm(width)

    for i, text in enumerate(_HEADER):
        if text is None:
            continue
        _write(tbl.cell(0, i), text, bold=True, size=8, center=True)

    tbl.cell(0, 4).merge(tbl.cell(0, 5))

    current = 1
    for ordinal, (emp, type_rows) in enumerate(employee_data, start=1):
        n = len(type_rows)
        start = current

        for i, (pt, run) in enumerate(type_rows):
            row = tbl.rows[current]

            if i == 0:
                _write(row.cells[0], str(ordinal), center=True)
                _write(row.cells[1], emp.high_risk_position_name or "")
                _write(row.cells[2],
                       f"{emp.first_name} {emp.last_name}".strip())

            interval = ""
            if run:
                b = run.process_binding
                interval = str(
                    b.custom_period_months
                    or b.process_type.default_period_months
                    or ""
                )
            _write(row.cells[3], interval, center=True)

            if pt:
                _write(row.cells[4], _short_label(pt.name))

            if run:
                rd = run.result_data or {}
                _write(row.cells[5], _fmt_date(run.performed_at))
                _write(row.cells[6], _fmt_date(run.valid_until))
                _write(row.cells[7], _result_data_field(
                    rd, "report_number", "broj_izvestaja"))
                _write(row.cells[8], _result_data_field(
                    rd, "fitness_assessment", "ocena_sposobnosti"))
                _write(row.cells[9], _result_data_field(
                    rd, "measures_taken", "preduzete_mere"))

            current += 1

        if n > 1:
            for col_idx in range(3):
                tbl.cell(start, col_idx).merge(
                    tbl.cell(start + n - 1, col_idx))

    doc.add_paragraph()

    sig = doc.add_table(rows=1, cols=2)
    sig.cell(0, 0).text = (
        "_____________________________\n"
        "Savetnik/saradnik za bezbednost i zdravlje na radu"
    )
    sig.cell(0, 1).text = "_____________________________\nPoslodavac"

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()
