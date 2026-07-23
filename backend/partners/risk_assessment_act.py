import io

import docx
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .kinney import category_label, is_high_risk
from .models import ClientCompany, JobRole


BASE_FONT = "Times New Roman"
HEADER_FILL = "D9D9D9"
INK = RGBColor(0x1A, 0x1A, 0x1A)


def _set_font(run, size=11, bold=False, italic=False, color=None, name=BASE_FONT):
    run.font.name = name
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:cs"), name)


def _set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _set_cell_margins(cell, top=40, bottom=40, left=80, right=80):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for edge, value in (("top", top), ("bottom", bottom),
                        ("start", left), ("end", right)):
        node = OxmlElement("w:" + edge)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    tc_pr.append(margins)


def _mark_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def _write(cell, text, bold=False, size=9, center=False, fill=None):
    cell.text = ""
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    _set_cell_margins(cell)
    if fill is not None:
        _set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(str(text) if text is not None else "")
    _set_font(r, size=size, bold=bold)


def _body(doc, text, size=11, italic=False, justify=True, space_after=8):
    p = doc.add_paragraph()
    if justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    _set_font(r, size=size, italic=italic)
    return p


def _heading(doc, number, text, size=13, space_before=14, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.keep_with_next = True
    label = f"{number} {text}" if number else text
    r = p.add_run(label)
    _set_font(r, size=size, bold=True, color=INK)
    return p


def _subheading(doc, number, text, size=11.5, space_before=12, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.keep_with_next = True
    label = f"{number} {text}" if number else text
    r = p.add_run(label)
    _set_font(r, size=size, bold=True, color=INK)
    return p


def _add_page_field(paragraph):
    run = paragraph.add_run()
    _set_font(run, size=9)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._element.append(begin)
    run._element.append(instr)
    run._element.append(end)


def _build_footer(section, company_name):
    footer = section.footer
    footer.is_linked_to_previous = False
    para = footer.paragraphs[0]
    para.text = ""
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    left = para.add_run(company_name + "\t")
    _set_font(left, size=9)
    tab_pr = para.paragraph_format.tab_stops
    try:
        from docx.enum.text import WD_TAB_ALIGNMENT
        section_width = section.page_width - section.left_margin - section.right_margin
        tab_pr.add_tab_stop(section_width, WD_TAB_ALIGNMENT.RIGHT)
    except Exception:
        pass
    label = para.add_run("Strana ")
    _set_font(label, size=9)
    _add_page_field(para)


_HAZARD_HEADER = [
    "Opasnost / štetnost",
    "V",
    "I",
    "P",
    "R",
    "Kategorija rizika",
    "Mere za otklanjanje / smanjenje",
]
_HAZARD_WIDTHS = [5.2, 0.9, 0.9, 0.9, 1.1, 3.3, 5.7]
_NUMERIC_COLS = {1, 2, 3, 4}


def _format_number(value):
    if value == int(value):
        return str(int(value))
    return str(value)


def _apply_column_widths(table, widths):
    table.autofit = False
    table.allow_autofit = False
    for col_idx, width in enumerate(widths):
        for cell in table.columns[col_idx].cells:
            cell.width = Cm(width)


def generate_risk_assessment_act(client_id: int) -> bytes:
    company = ClientCompany.objects.get(pk=client_id)
    roles = (
        JobRole.objects.filter(client_company_id=client_id)
        .prefetch_related("hazards", "hazards__hazard")
        .order_by("name")
    )

    doc = docx.Document()

    normal = doc.styles["Normal"]
    normal.font.name = BASE_FONT
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal_rpr = normal.element.get_or_add_rPr()
    normal_rfonts = normal_rpr.find(qn("w:rFonts"))
    if normal_rfonts is None:
        normal_rfonts = OxmlElement("w:rFonts")
        normal_rpr.append(normal_rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        normal_rfonts.set(qn(attr), BASE_FONT)

    section = doc.sections[0]
    section.left_margin = section.right_margin = Cm(2.0)
    section.top_margin = section.bottom_margin = Cm(2.0)

    _build_footer(section, company.name)

    for _ in range(6):
        doc.add_paragraph()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(4)
    t = title.add_run("AKT O PROCENI RIZIKA")
    _set_font(t, size=28, bold=True, color=INK)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(4)
    st = subtitle.add_run("na radnom mestu i u radnoj okolini")
    _set_font(st, size=13, italic=True)

    for _ in range(3):
        doc.add_paragraph()

    identity = [
        ("Poslovno ime poslodavca", company.name or ""),
        ("Adresa sedišta poslodavca", company.address or ""),
        ("PIB poslodavca", company.tax_id or ""),
    ]
    ident_tbl = doc.add_table(rows=len(identity), cols=2)
    ident_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    _apply_column_widths(ident_tbl, [6.0, 9.0])
    for idx, (label, value) in enumerate(identity):
        _write(ident_tbl.cell(idx, 0), label, bold=True, size=11,
               fill=HEADER_FILL)
        _write(ident_tbl.cell(idx, 1), value, size=11)
    ident_tbl.style = "Table Grid"

    for _ in range(6):
        doc.add_paragraph()

    place_date = doc.add_paragraph()
    place_date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pd = place_date.add_run("Mesto i datum izrade: ______________________")
    _set_font(pd, size=11)

    new_section = doc.add_section(WD_SECTION.NEW_PAGE)
    new_section.left_margin = new_section.right_margin = Cm(2.0)
    new_section.top_margin = new_section.bottom_margin = Cm(2.0)
    _build_footer(new_section, company.name)

    _heading(doc, "1.", "ODLUKA O POKRETANJU POSTUPKA PROCENE RIZIKA")
    _body(
        doc,
        f"Poslodavac {company.name} pokreće postupak procene rizika na svim "
        "radnim mestima i u radnoj okolini, u skladu sa Zakonom o bezbednosti "
        "i zdravlju na radu („Sl. glasnik RS”, br. 35/2023) i Pravilnikom o "
        "načinu i postupku procene rizika na radnom mestu i u radnoj okolini.",
    )

    _heading(doc, "2.", "UVOD")
    _body(
        doc,
        "Za svako radno mesto utvrđene su opasnosti i štetnosti, procenjen je "
        "nivo rizika i određene su mere za otklanjanje ili smanjenje rizika.",
    )

    _heading(doc, "3.", "PROCENA RIZIKA PO RADNIM MESTIMA")

    high_risk_roles = []
    role_number = 0
    for role in roles:
        role_number += 1
        hazards = list(role.hazards.all())
        _subheading(doc, f"3.{role_number}.", f"Radno mesto: {role.name}")
        if not hazards:
            _body(
                doc,
                "Za ovo radno mesto nisu unete opasnosti i štetnosti.",
                italic=True,
            )
            continue

        tbl = doc.add_table(rows=1 + len(hazards), cols=len(_HAZARD_HEADER))
        tbl.style = "Table Grid"
        _apply_column_widths(tbl, _HAZARD_WIDTHS)

        for i, text in enumerate(_HAZARD_HEADER):
            _write(tbl.cell(0, i), text, bold=True, size=9, center=True,
                   fill=HEADER_FILL)
        _mark_header_row(tbl.rows[0])

        role_is_high = False
        for idx, jrh in enumerate(hazards, start=1):
            row = tbl.rows[idx]
            _write(row.cells[0], jrh.hazard.label, size=9)
            _write(row.cells[1], _format_number(jrh.verovatnoca),
                   size=9, center=True)
            _write(row.cells[2], _format_number(jrh.izlozenost),
                   size=9, center=True)
            _write(row.cells[3], _format_number(jrh.posledica),
                   size=9, center=True)
            _write(row.cells[4], _format_number(jrh.rizik),
                   size=9, center=True, bold=True)
            _write(row.cells[5], category_label(jrh.risk_category), size=9,
                   center=True)
            _write(row.cells[6], jrh.mere or "", size=9)
            if is_high_risk(jrh.risk_category):
                role_is_high = True

        if role_is_high:
            high_risk_roles.append(role.name)

        doc.add_paragraph()

    _heading(doc, "4.", "ZAKLJUČAK")
    if high_risk_roles:
        _body(
            doc,
            "Na osnovu sprovedene procene rizika, radnim mestima sa povećanim "
            "rizikom utvrđena su sledeća radna mesta:",
        )
        for name in high_risk_roles:
            p = doc.add_paragraph(style="List Bullet")
            r = p.add_run(name)
            _set_font(r, size=11)
    else:
        _body(
            doc,
            "Na osnovu sprovedene procene rizika, nije utvrđeno nijedno radno "
            "mesto sa povećanim rizikom.",
        )

    _heading(doc, "5.", "IZJAVA POSLODAVCA")
    _body(
        doc,
        f"Poslodavac {company.name} izjavljuje da će sprovoditi utvrđene mere "
        "za bezbednost i zdravlje na radu, pratiti njihovu primenu i vršiti "
        "izmene i dopune ovog akta u slučaju pojave novih opasnosti i "
        "štetnosti ili promene nivoa rizika na radnom mestu.",
    )

    for _ in range(3):
        doc.add_paragraph()

    sig = doc.add_table(rows=1, cols=2)
    sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    _apply_column_widths(sig, [7.5, 7.5])
    _write(
        sig.cell(0, 0),
        "_____________________________\nLice za bezbednost i zdravlje na radu",
        size=11,
        center=True,
    )
    _write(
        sig.cell(0, 1),
        "_____________________________\nPoslodavac",
        size=11,
        center=True,
    )

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()
