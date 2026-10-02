import io
import re

import docx
from django.test import SimpleTestCase

from documents.word_engine import EACH_RE, lookup, render_docx, template_tags
from partners.document_contexts import preview_context
from partners.document_service import MASTER_FILES, master_path


def _all_text(content):
    document = docx.Document(io.BytesIO(content))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    for section in document.sections:
        parts += [p.text for p in section.header.paragraphs]
        parts += [p.text for p in section.footer.paragraphs]
    return "\n".join(parts)


def _first(value):
    if isinstance(value, list):
        return value[0] if value else None
    return value


def _resolve(path, context, aliases):
    head, _, rest = path.partition(".")
    if head in aliases:
        current = aliases[head]
        parts = rest.split(".") if rest else []
    else:
        current = context
        parts = path.split(".")
    for part in parts:
        current = _first(current)
        if not isinstance(current, dict) or part not in current:
            return False
        current = current[part]
    return True


class DocumentMastersTest(SimpleTestCase):
    def test_every_master_renders_without_leftover_tags(self):
        context = preview_context()
        for code in MASTER_FILES:
            with self.subTest(code=code):
                content = render_docx(master_path(code).read_bytes(), context)
                self.assertNotIn("{{", _all_text(content))

    def test_every_master_tag_exists_in_context(self):
        context = preview_context()
        for code in MASTER_FILES:
            source = master_path(code).read_bytes()
            text = _all_text(source)
            aliases = {}
            for match in re.finditer(r"\{\{\s*#each[^}]*\}\}", text):
                each = EACH_RE.match(match.group(0))
                if each:
                    aliases[each.group(2)] = _first(lookup(each.group(1), context)) or {}
            for tag in template_tags(source):
                with self.subTest(code=code, tag=tag):
                    self.assertTrue(
                        _resolve(tag, context, aliases),
                        f"{code}: {tag} is not provided by the document context",
                    )
