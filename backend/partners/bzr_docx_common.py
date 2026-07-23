import io

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt


def new_doc():
    doc = docx.Document()
    section = doc.sections[0]
    section.left_margin = section.right_margin = Cm(2.0)
    section.top_margin = section.bottom_margin = Cm(2.0)
    return doc


def title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(14)
    return p


def subtitle(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.font.size = Pt(11)
    return p


def heading(doc, text, size=12):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(size)
    return p


def para(doc, text, bold=False, size=11):
    p = doc.add_paragraph()
    r = p.add_run(text if text is not None else "")
    r.bold = bold
    r.font.size = Pt(size)
    return p


def bullet(doc, text):
    return doc.add_paragraph(text, style="List Bullet")


def numbered(doc, text):
    return doc.add_paragraph(text, style="List Number")


def spacer(doc):
    doc.add_paragraph()


def _cell(cell, text, bold=False, size=10):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(str(text) if text is not None else "")
    r.bold = bold
    r.font.size = Pt(size)


def company_header(doc, company):
    tbl = doc.add_table(rows=2, cols=3)
    tbl.style = "Table Grid"
    _cell(tbl.cell(0, 0), company.name, bold=True)
    _cell(tbl.cell(0, 1), company.address or "")
    _cell(tbl.cell(0, 2), company.tax_id or "")
    _cell(tbl.cell(1, 0), "Poslovno ime poslodavca", size=8)
    _cell(tbl.cell(1, 1), "Adresa sedišta", size=8)
    _cell(tbl.cell(1, 2), "PIB", size=8)
    return tbl


def signatures(doc, left="Lice za bezbednost i zdravlje na radu",
               right="Poslodavac"):
    spacer(doc)
    tbl = doc.add_table(rows=1, cols=2)
    tbl.cell(0, 0).text = f"_____________________________\n{left}"
    tbl.cell(0, 1).text = f"_____________________________\n{right}"
    return tbl


def director_name(company):
    try:
        persons = list(company.contact_persons.all())
    except Exception:
        persons = []
    directors = [p for p in persons if p.role == "DIRECTOR"]
    d = next((p for p in directors if p.is_primary), None) or (
        directors[0] if directors else None)
    return getattr(d, "full_name", "") or ""


def finalize(doc):
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()
