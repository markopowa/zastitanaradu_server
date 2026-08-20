import io
import uuid
from datetime import date

import docx
from django.contrib.auth import get_user_model
from django.core.files.base import ContentFile
from django.test import TestCase
from rest_framework.test import APIClient

from documents.models import DocumentCategory, DocumentFile, DocumentTemplate
from partners.document_generation import generate_employee_document
from partners.models import (
    ClientCompany,
    Employee,
    Hazard,
    JobRole,
    JobRoleHazard,
    JobRoleLZO,
)
from partners.obrazac6 import build_obrazac6_data
from processes.models import ProcessBinding, ProcessRun, ProcessRunDocument, ProcessType

User = get_user_model()


def make_company(**kwargs):
    defaults = {"name": "Test firma", "tax_id": uuid.uuid4().hex[:9]}
    defaults.update(kwargs)
    return ClientCompany.objects.create(**defaults)


def make_process_type(code, **kwargs):
    defaults = {
        "code": code,
        "name": f"Proces {code}",
        "subject_kind": ProcessType.SUBJECT_EMPLOYEE,
        "domain": ProcessType.DOMAIN_BZNR,
        "shape": ProcessType.SHAPE_PERIODIC,
        "proof_kind": ProcessType.PROOF_RECORDED,
        "applicability_rule": {"always": True},
        "is_active": True,
    }
    defaults.update(kwargs)
    return ProcessType.objects.create(**defaults)


def build_cell_docx() -> bytes:
    doc = docx.Document()
    table = doc.add_table(rows=1, cols=3)
    table.cell(0, 0).text = "label"
    table.cell(0, 1).text = ""
    table.cell(0, 2).text = "other"
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


def build_obrazac6_blank_docx() -> bytes:
    doc = docx.Document()
    labels = [
        "Naziv radnog mesta",
        "Opis poslova na tom radnom mestu",
        "Slučaj, odnosno razlog izvršene obuke za bezbedan i zdrav rad",
        "Naziv lične zaštitne opreme (ako je korišćenje LZO utvrđeno aktom o proceni rizika)",
        "",
        "",
        "",
        "Opasnosti, odnosno štetnosti sa kojima je zaposleni upoznat",
        "Konkretne mere za bezbedan i zdrav rad",
    ]
    table = doc.add_table(rows=len(labels), cols=3)
    for i, label in enumerate(labels):
        table.rows[i].cells[0].text = label
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


def read_cell(content_bytes: bytes, table_index: int, row: int, col: int) -> str:
    doc = docx.Document(io.BytesIO(content_bytes))
    return doc.tables[table_index].rows[row].cells[col].text


def all_text(content_bytes: bytes) -> str:
    doc = docx.Document(io.BytesIO(content_bytes))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


class OnDemandGenerationTestCase(TestCase):
    def setUp(self):
        self.user = User.objects.create_superuser(
            username="ondemandgen", password="x", email="ondemandgen@test.local",
        )
        self.category = DocumentCategory.objects.create(name="Test kategorija")
        self.obrazac6_doc_tpl = DocumentTemplate.objects.create(
            name="Obrazac 6 — evidencija o osposobljenosti za bezbedan rad",
            context_type=DocumentTemplate.CONTEXT_EMPLOYEE,
            category=self.category,
            generation_config={"mode": "VISUAL", "placeholders": []},
        )
        self.obrazac6_doc_tpl.template_file.save(
            "obrazac6_blank.docx",
            ContentFile(build_obrazac6_blank_docx()),
            save=True,
        )
        self.lzo_doc_tpl = DocumentTemplate.objects.create(
            name="Karton zaduženja LZO (revers)",
            context_type=DocumentTemplate.CONTEXT_EMPLOYEE,
            category=self.category,
            generation_config={},
        )
        self.obrazac6_process_type = make_process_type("OSPOSOBLJAVANJE_BZR")
        self.lzo_process_type = make_process_type("LZO_ZADUZENJE")
        self.company = make_company(name="Firma Test")

    def _make_employee_with_role(self, first_name, last_name, with_data=True):
        role = JobRole.objects.create(
            client_company=self.company,
            name=f"Radno mesto {first_name}",
            description="Poslovi na radnom mestu.",
        )
        if with_data:
            hazard = Hazard.objects.create(
                code=f"HZ-{uuid.uuid4().hex[:8]}",
                label="Mehaničke opasnosti",
                kind=Hazard.KIND_OPASNOST,
            )
            JobRoleHazard.objects.create(
                job_role=role,
                hazard=hazard,
                verovatnoca=6,
                izlozenost=6,
                posledica=15,
                mere="Obuka, LZO, redovan pregled.",
            )
            JobRoleLZO.objects.create(
                job_role=role,
                name="Zaštitne cipele sa čeličnom kapom",
                standard="SRPS EN ISO 20345",
                interval_months=12,
            )
        return Employee.objects.create(
            client_company=self.company,
            first_name=first_name,
            last_name=last_name,
            job_role=role,
        )

    def _make_binding_and_run(self, employee, process_type):
        binding = ProcessBinding.objects.create(
            process_type=process_type,
            subject_kind=ProcessBinding.SUBJECT_EMPLOYEE,
            employee=employee,
            is_active=True,
            next_run_at=date.today(),
        )
        run = ProcessRun.objects.create(
            process_binding=binding,
            process_type=process_type,
            scheduled_for=date.today(),
            status=ProcessRun.STATUS_PENDING,
        )
        return binding, run

    def test_obrazac6_data_is_data_driven_from_role(self):
        employee = self._make_employee_with_role("Marko", "Markovic")
        data = build_obrazac6_data(employee)

        self.assertEqual(data["role_name"], "Radno mesto Marko")
        self.assertIn("Mehaničke opasnosti", data["opasnosti"])
        self.assertIn("Obuka, LZO", data["mere"])
        self.assertEqual(
            [item.name for item in data["lzo"]],
            ["Zaštitne cipele sa čeličnom kapom"],
        )

    def test_service_generates_obrazac6_docx_from_role_data(self):
        employee = self._make_employee_with_role("Marko", "Markovic")
        self._make_binding_and_run(employee, self.obrazac6_process_type)

        doc_file, content_bytes = generate_employee_document(employee, "OBRAZAC6")

        self.assertTrue(content_bytes.startswith(b"%PDF"))
        self.assertTrue(doc_file.file.name.endswith(".docx"))
        self.assertTrue(
            ProcessRunDocument.objects.filter(document_file=doc_file).exists()
        )

        doc_file_2, _ = generate_employee_document(employee, "OBRAZAC6")
        self.assertNotEqual(doc_file.id, doc_file_2.id)
        self.assertEqual(
            ProcessRunDocument.objects.filter(
                process_run__process_binding__employee=employee,
            ).count(),
            2,
        )

    def test_service_generates_lzo_revers_from_role_data(self):
        employee = self._make_employee_with_role("Ana", "Anic")
        self._make_binding_and_run(employee, self.lzo_process_type)

        doc_file, content_bytes = generate_employee_document(employee, "LZO_REVERS")

        text = all_text(content_bytes)
        self.assertIn("Ana Anic", text)
        self.assertIn("Zaštitne cipele sa čeličnom kapom", text)
        self.assertTrue(
            ProcessRunDocument.objects.filter(document_file=doc_file).exists()
        )

    def test_no_obligation_raises_value_error(self):
        employee = self._make_employee_with_role("Petar", "Petrovic")
        with self.assertRaises(ValueError) as ctx:
            generate_employee_document(employee, "OBRAZAC6")
        self.assertIn("nema obavezu", str(ctx.exception))

    def test_no_obligation_returns_400_via_endpoint(self):
        employee = self._make_employee_with_role("Petar", "Petrovic")
        client = APIClient()
        client.force_authenticate(user=self.user)
        response = client.post(
            f"/api/partners/employees/{employee.id}/generate-document/",
            {"kind": "OBRAZAC6"},
            format="json",
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("nema obavezu", response.data["detail"])

    def test_endpoint_returns_pdf_download(self):
        employee = self._make_employee_with_role("Jovan", "Jovanovic")
        self._make_binding_and_run(employee, self.obrazac6_process_type)
        client = APIClient()
        client.force_authenticate(user=self.user)
        response = client.post(
            f"/api/partners/employees/{employee.id}/generate-document/",
            {"kind": "OBRAZAC6"},
            format="json",
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("attachment", response["Content-Disposition"])
        self.assertTrue(response.content.startswith(b"%PDF"))

    def test_company_generate_all_merges_pdf_and_reports_counts(self):
        emp1 = self._make_employee_with_role("Prvi", "Zaposleni")
        emp2 = self._make_employee_with_role("Drugi", "Zaposleni")
        self._make_employee_with_role("Treci", "Zaposleni")
        self._make_binding_and_run(emp1, self.obrazac6_process_type)
        self._make_binding_and_run(emp2, self.obrazac6_process_type)

        before_count = DocumentFile.objects.count()

        client = APIClient()
        client.force_authenticate(user=self.user)
        response = client.post(
            f"/api/partners/client-companies/{self.company.id}/generate-obrazac6-all/",
        )

        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.content.startswith(b"%PDF"))
        self.assertEqual(response["X-Generated-Count"], "2")
        self.assertEqual(response["X-Skipped-Count"], "1")
        self.assertEqual(DocumentFile.objects.count(), before_count + 2)
