import io
from copy import deepcopy

import docx
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

CELL_MAP = [
    ("Ime i prezime zaposlenog", "employee_name"),
    ("Naziv radnog mesta", "role_name"),
    ("Opis poslova na tom radnom mestu", "role_description"),
    ("Slučaj, odnosno razlog", "razlog"),
    ("Opasnosti, odnosno štetnosti", "opasnosti"),
    ("Konkretne mere", "mere"),
]


def build_obrazac6_data(employee, context=None):
    context = context or {}
    role = getattr(employee, "job_role", None)
    hazards = list(role.hazards.select_related("hazard").all()) if role else []
    opasnosti = "\n".join(
        f"{i}. {h.hazard.label}" for i, h in enumerate(hazards, 1)
    )
    mere = "\n".join(
        f"{i}. {h.mere}"
        for i, h in enumerate(hazards, 1)
        if (h.mere or "").strip()
    )
    lzo = list(role.lzo_items.all()) if role else []
    emp_ctx = context.get("employee") or {}
    full_name = (emp_ctx.get("full_name") or "").strip()
    if not full_name:
        first = getattr(employee, "first_name", "") or ""
        last = getattr(employee, "last_name", "") or ""
        full_name = f"{first} {last}".strip()
    performed = (context.get("performed_at") or "").strip()
    if not performed:
        performed = (context.get("scheduled_for") or "").strip()
    return {
        "employee_name": full_name,
        "role_name": role.name if role else "",
        "role_description": (
            (getattr(role, "description", "") or "") if role else ""
        ),
        "razlog": "01",
        "opasnosti": opasnosti,
        "mere": mere,
        "lzo": lzo,
        "performed_at": performed,
    }


def _unique_cells(row):
    seen = set()
    out = []
    for cell in row.cells:
        key = id(cell._tc)
        if key in seen:
            continue
        seen.add(key)
        out.append(cell)
    return out


def _set_cell(cell, text):
    for p in list(cell.paragraphs):
        for r in list(p.runs):
            r._element.getparent().remove(r._element)
    lines = (text or "").split("\n")
    p0 = cell.paragraphs[0]
    p0.add_run(lines[0] if lines else "")
    for line in lines[1:]:
        para = cell.add_paragraph()
        para.add_run(line)


def _fill_value_cells(row, text):
    cells = _unique_cells(row)
    if len(cells) < 2:
        return
    for cell in cells[1:]:
        _set_cell(cell, text)


def _fill_date_row(table, performed_at):
    if not performed_at:
        return
    rows = table.rows
    for i, row in enumerate(rows):
        cells = _unique_cells(row)
        if not cells:
            continue
        label = (cells[0].text or "").strip()
        if not label.startswith("teorijske"):
            continue
        if i + 1 >= len(rows):
            return
        for cell in _unique_cells(rows[i + 1]):
            _set_cell(cell, performed_at)
        return


def _fill_lzo_rows(table, lzo):
    if not lzo:
        return
    rows = table.rows
    start = end = None
    for i, row in enumerate(rows):
        cells = _unique_cells(row)
        label = (cells[0].text or "").strip() if cells else ""
        if label.startswith("Naziv lične zaštitne opreme"):
            start = i + 1
        elif start is not None and label.startswith("Opasnosti"):
            end = i
            break
    if start is None or end is None:
        return
    names = [
        (getattr(item, "name", "") or "").strip()
        for item in lzo
        if (getattr(item, "name", "") or "").strip()
    ]
    if not names:
        return
    if end <= start:
        ref = table.rows[end]._tr
        new_tr = deepcopy(ref)
        for tc in new_tr.iterchildren():
            if not tc.tag.endswith("}tc"):
                continue
            for p in list(tc.iterchildren()):
                if p.tag.endswith("}p"):
                    for child in list(p):
                        if child.tag.endswith("}r"):
                            p.remove(child)
        ref.addprevious(new_tr)
        end = start + 1
        rows = table.rows
    cells = _unique_cells(rows[start])
    if cells:
        _set_cell(cells[0], "\n".join(names))


def _set_paragraph_text(para, text):
    for r in list(para.runs):
        r._element.getparent().remove(r._element)
    para.add_run(text)


def _insert_paragraph_after(paragraph, text: str):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    new_para.add_run(text)
    return new_para


def _fill_header(doc, context):
    client = context.get("client") or {}
    name = (client.get("name") or "").strip()
    address = (client.get("address") or "").strip()
    tax_id = (client.get("tax_id") or "").strip()
    if not (name or address or tax_id):
        return
    values_line = "\t".join([name, address, tax_id])
    for para in doc.paragraphs:
        text = para.text or ""
        if "Poslovno ime poslodavca" not in text:
            continue
        nxt = para._p.getnext()
        if nxt is not None and nxt.tag.endswith("}p"):
            nxt_para = Paragraph(nxt, para._parent)
            if not (nxt_para.text or "").strip():
                _set_paragraph_text(nxt_para, values_line)
                return
        _insert_paragraph_after(para, values_line)
        return


def _fill_cells(doc, data):
    for table in doc.tables:
        for row in table.rows:
            cells = _unique_cells(row)
            if not cells:
                continue
            label = (cells[0].text or "").strip()
            for prefix, key in CELL_MAP:
                if label.startswith(prefix):
                    _fill_value_cells(row, data.get(key) or "")
                    break
        _fill_lzo_rows(table, data.get("lzo") or [])
        _fill_date_row(table, data.get("performed_at") or "")


def generate_obrazac6(employee, context, template_file, placements=None) -> bytes:
    data = build_obrazac6_data(employee, context)
    with template_file.open("rb") as fh:
        doc = docx.Document(io.BytesIO(fh.read()))
    _fill_header(doc, context or {})
    _fill_cells(doc, data)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
