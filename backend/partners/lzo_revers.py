import io

import docx


def _set_paragraph_text(para, text):
    for r in list(para.runs):
        r._element.getparent().remove(r._element)
    para.add_run(text)


def _fill_labeled_paragraphs(doc, labels_to_values):
    for para in doc.paragraphs:
        text = (para.text or "").strip()
        for label, value in labels_to_values.items():
            if text.startswith(label):
                suffix = (value or "").strip()
                new_text = f"{label} {suffix}".rstrip() if suffix else label
                _set_paragraph_text(para, new_text)
                break


def _fill_lzo_table(doc, lzo_items):
    if not doc.tables:
        return
    table = doc.tables[0]
    while len(table.rows) > 1:
        tr = table.rows[-1]._tr
        tr.getparent().remove(tr)
    for index, item in enumerate(lzo_items, start=1):
        row = table.add_row()
        cells = row.cells
        values = [
            str(index),
            getattr(item, "name", "") or "",
            getattr(item, "standard", "") or "",
            (
                str(item.interval_months)
                if getattr(item, "interval_months", None) is not None
                else ""
            ),
            "",
            "",
            "",
            "",
        ]
        for cell, value in zip(cells, values):
            cell.text = value


def generate_lzo_revers(employee, context=None, template_file=None) -> bytes:
    context = context or {}
    if template_file is None:
        from documents.models import DocumentTemplate

        doc_template = DocumentTemplate.objects.filter(
            name="Karton zaduženja LZO (revers)",
        ).first()
        if not doc_template or not doc_template.template_file:
            raise ValueError("Šablon „Karton zaduženja LZO (revers)” nije podešen.")
        template_file = doc_template.template_file

    with template_file.open("rb") as fh:
        doc = docx.Document(io.BytesIO(fh.read()))

    emp_ctx = context.get("employee") or {}
    client_ctx = context.get("client") or {}
    first = getattr(employee, "first_name", "") or ""
    last = getattr(employee, "last_name", "") or ""
    full_name = (emp_ctx.get("full_name") or f"{first} {last}").strip()
    role = getattr(employee, "job_role", None)
    role_name = (
        emp_ctx.get("position")
        or (getattr(role, "name", "") if role is not None else "")
        or ""
    )
    company_name = (client_ctx.get("name") or "").strip()
    if not company_name:
        company = getattr(employee, "client_company", None)
        company_name = getattr(company, "name", "") if company is not None else ""

    _fill_labeled_paragraphs(
        doc,
        {
            "Ime i prezime zaposlenog:": full_name,
            "Naziv radnog mesta:": role_name,
            "Naziv firme:": company_name,
        },
    )

    lzo_items = list(role.lzo_items.all()) if role is not None else []
    _fill_lzo_table(doc, lzo_items)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
