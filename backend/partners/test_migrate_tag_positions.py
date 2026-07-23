import json
import shutil
import tempfile
import uuid
from pathlib import Path

import docx
from django.conf import settings
from django.core.files.base import ContentFile
from django.core.management import call_command
from django.test import TestCase, override_settings

from documents.models import DocumentCategory, DocumentTemplate
from partners.management.commands.add_setup import Command as AddSetupCommand
from partners.management.commands.add_setup import MASTER_TEMPLATE_FILES
from partners.models import ClientCompany, JobRole


def build_docx_bytes(build_fn) -> bytes:
    import io
    doc = docx.Document()
    build_fn(doc)
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


def require_libreoffice(test_case) -> None:
    if not shutil.which(settings.LIBREOFFICE_BIN):
        test_case.skipTest("LibreOffice not available on this host")


class MigrateTagPositionsSkipTests(TestCase):
    """No-tags files must be left untouched — this needs no LibreOffice."""

    def setUp(self):
        self.category = DocumentCategory.objects.create(
            name=f"Kategorija {uuid.uuid4().hex[:8]}")

    def test_master_without_tags_is_skipped(self):
        tpl = DocumentTemplate.objects.create(
            name=f"Sablon {uuid.uuid4().hex[:8]}",
            context_type=DocumentTemplate.CONTEXT_EMPLOYEE,
            category=self.category,
            generation_config={"mode": "DOCX_PLACEHOLDER"},
        )
        tpl.template_file.save(
            "no_tags.docx",
            ContentFile(build_docx_bytes(
                lambda d: d.add_paragraph("Bez tagova."))),
            save=True,
        )
        call_command("migrate_tag_positions")
        tpl.refresh_from_db()
        self.assertEqual(tpl.generation_config, {"mode": "DOCX_PLACEHOLDER"})

    def test_role_blank_without_tags_is_skipped(self):
        company = ClientCompany.objects.create(
            name="Firma", tax_id=uuid.uuid4().hex[:9])
        role = JobRole.objects.create(client_company=company, name="Radnik")
        role.obrazac6_template.save(
            "no_tags.docx",
            ContentFile(build_docx_bytes(
                lambda d: d.add_paragraph("Bez tagova."))),
            save=True,
        )
        call_command("migrate_tag_positions")
        role.refresh_from_db()
        self.assertEqual(role.obrazac6_fields, [])


class MigrateTagPositionsDbModeTests(TestCase):
    def setUp(self):
        self.category = DocumentCategory.objects.create(
            name=f"Kategorija {uuid.uuid4().hex[:8]}")

    def test_master_with_tags_becomes_visual_and_tags_are_stripped(self):
        require_libreoffice(self)
        tpl = DocumentTemplate.objects.create(
            name=f"Sablon {uuid.uuid4().hex[:8]}",
            context_type=DocumentTemplate.CONTEXT_EMPLOYEE,
            category=self.category,
            generation_config={"mode": "DOCX_PLACEHOLDER"},
        )
        tpl.template_file.save(
            "with_tags.docx",
            ContentFile(build_docx_bytes(
                lambda d: d.add_paragraph(
                    "Zaposleni: {{ employee.full_name }}"))),
            save=True,
        )

        call_command("migrate_tag_positions")

        tpl.refresh_from_db()
        self.assertEqual(tpl.generation_config.get("mode"), "VISUAL")
        placeholders = tpl.generation_config.get("placeholders") or []
        self.assertEqual(len(placeholders), 1)
        self.assertEqual(placeholders[0]["fieldKey"], "employee.full_name")

        with tpl.template_file.open("rb") as fh:
            saved_doc = docx.Document(fh)
        all_text = "\n".join(p.text for p in saved_doc.paragraphs)
        self.assertNotIn("{{", all_text)

        # Idempotent: a second run makes no further changes.
        call_command("migrate_tag_positions")
        tpl.refresh_from_db()
        self.assertEqual(
            tpl.generation_config.get("placeholders"), placeholders)

    def test_dry_run_does_not_write_anything(self):
        require_libreoffice(self)
        tpl = DocumentTemplate.objects.create(
            name=f"Sablon {uuid.uuid4().hex[:8]}",
            context_type=DocumentTemplate.CONTEXT_EMPLOYEE,
            category=self.category,
            generation_config={"mode": "DOCX_PLACEHOLDER"},
        )
        tpl.template_file.save(
            "with_tags.docx",
            ContentFile(build_docx_bytes(
                lambda d: d.add_paragraph(
                    "Zaposleni: {{ employee.full_name }}"))),
            save=True,
        )

        call_command("migrate_tag_positions", "--dry-run")

        tpl.refresh_from_db()
        self.assertEqual(tpl.generation_config, {"mode": "DOCX_PLACEHOLDER"})
        with tpl.template_file.open("rb") as fh:
            saved_doc = docx.Document(fh)
        all_text = "\n".join(p.text for p in saved_doc.paragraphs)
        self.assertIn("{{", all_text)


class MigrateTagPositionsFilesOnlyModeTests(TestCase):
    def setUp(self):
        self.tmp_dir = Path(tempfile.mkdtemp(prefix="test_files_only_"))
        self.addCleanup(shutil.rmtree, self.tmp_dir, ignore_errors=True)
        self.master_dir = self.tmp_dir / "master_templates"
        self.master_dir.mkdir()
        self.blanks_dir = self.tmp_dir / "pilot_blanks" / "obrazac6"
        self.blanks_dir.mkdir(parents=True)

    def _write_docx(self, path: Path, build_fn) -> None:
        doc = docx.Document()
        build_fn(doc)
        doc.save(str(path))

    def test_files_only_writes_sidecars_and_strips_tags(self):
        require_libreoffice(self)
        master_name, master_filename = next(iter(MASTER_TEMPLATE_FILES.items()))
        master_path = self.master_dir / master_filename
        self._write_docx(
            master_path,
            lambda d: d.add_paragraph("Firma: {{ client.name }}"),
        )
        blank_path = self.blanks_dir / "obrazac6_test_role.docx"
        self._write_docx(
            blank_path,
            lambda d: d.add_paragraph("Zaposleni: {{ employee.full_name }}"),
        )

        with override_settings(BASE_DIR=self.tmp_dir):
            call_command("migrate_tag_positions", "--files-only")

        master_sidecar = json.loads(
            (self.master_dir / "placements.json").read_text())
        self.assertIn(master_name, master_sidecar)
        self.assertEqual(master_sidecar[master_name]["mode"], "VISUAL")
        self.assertEqual(
            len(master_sidecar[master_name]["placeholders"]), 1)

        blank_sidecar = json.loads(
            (self.blanks_dir / "placements.json").read_text())
        self.assertIn("obrazac6_test_role.docx", blank_sidecar)
        self.assertEqual(len(blank_sidecar["obrazac6_test_role.docx"]), 1)

        master_text = "\n".join(
            p.text for p in docx.Document(str(master_path)).paragraphs)
        self.assertNotIn("{{", master_text)
        blank_text = "\n".join(
            p.text for p in docx.Document(str(blank_path)).paragraphs)
        self.assertNotIn("{{", blank_text)

        # Idempotent second run: already tag-free, sidecars preserved as-is.
        with override_settings(BASE_DIR=self.tmp_dir):
            call_command("migrate_tag_positions", "--files-only")
        master_sidecar_after = json.loads(
            (self.master_dir / "placements.json").read_text())
        self.assertEqual(master_sidecar_after, master_sidecar)


class SidecarSeedingTests(TestCase):
    """add_setup and attach_obrazac6_blanks must consume placements.json
    sidecars without needing LibreOffice — the sidecar is just JSON."""

    def setUp(self):
        self.tmp_dir = Path(tempfile.mkdtemp(prefix="test_sidecar_seed_"))
        self.addCleanup(shutil.rmtree, self.tmp_dir, ignore_errors=True)

    def test_add_setup_seeds_generation_config_from_sidecar(self):
        master_name, master_filename = next(iter(MASTER_TEMPLATE_FILES.items()))
        master_dir = self.tmp_dir / "master_templates"
        master_dir.mkdir()
        (master_dir / master_filename).write_bytes(
            build_docx_bytes(lambda d: d.add_paragraph("Firma: ")),
        )
        derived_config = {
            "mode": "VISUAL",
            "placeholders": [{
                "fieldKey": "client.name",
                "page": 0, "xPct": 10, "yPct": 10,
                "widthPct": 30, "heightPct": 5, "fontSize": 11,
            }],
        }
        (master_dir / "placements.json").write_text(
            json.dumps({master_name: derived_config}))

        category = DocumentCategory.objects.create(
            name=f"Kategorija {uuid.uuid4().hex[:8]}")
        tpl = DocumentTemplate.objects.create(
            name=master_name,
            context_type=DocumentTemplate.CONTEXT_CLIENT_COMPANY,
            category=category,
        )

        with override_settings(BASE_DIR=self.tmp_dir):
            AddSetupCommand()._attach_master_templates()

        tpl.refresh_from_db()
        self.assertEqual(tpl.generation_config, derived_config)

    def test_attach_obrazac6_blanks_seeds_role_fields_from_sidecar(self):
        blanks_dir = self.tmp_dir / "pilot_blanks" / "obrazac6"
        blanks_dir.mkdir(parents=True)
        company = ClientCompany.objects.create(
            name="Firma", tax_id=uuid.uuid4().hex[:9])
        role = JobRole.objects.create(client_company=company, name="Varilac")
        filename = "obrazac6_varilac.docx"
        (blanks_dir / filename).write_bytes(
            build_docx_bytes(lambda d: d.add_paragraph("Zaposleni: ")),
        )
        placements = [{
            "fieldKey": "employee.full_name",
            "page": 0, "xPct": 5, "yPct": 5,
            "widthPct": 20, "heightPct": 4, "fontSize": 10,
        }]
        (blanks_dir / "placements.json").write_text(
            json.dumps({filename: placements}))

        with override_settings(BASE_DIR=self.tmp_dir):
            call_command("attach_obrazac6_blanks", "--kind", "obrazac6")

        role.refresh_from_db()
        self.assertTrue(bool(role.obrazac6_template))
        self.assertEqual(role.obrazac6_fields, placements)
