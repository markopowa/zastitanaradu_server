import io
import uuid

import docx
from django.contrib.auth import get_user_model
from rest_framework.test import APITestCase

from partners.bzr_documents import bzr_document_catalog, generate_bzr_document
from partners.models import ClientCompany, ContactPerson

User = get_user_model()


def doc_text(content):
    d = docx.Document(io.BytesIO(content))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


class BzrDocumentsUnitTest(APITestCase):
    def setUp(self):
        self.company = ClientCompany.objects.create(
            name="Pekara Klas", tax_id=uuid.uuid4().hex[:9],
            address="Ulica 1, Beograd")
        ContactPerson.objects.create(
            client_company=self.company, full_name="Marko Marković",
            role="DIRECTOR", is_primary=True)

    def test_every_catalog_kind_generates_valid_docx_without_dashes(self):
        catalog = bzr_document_catalog()
        self.assertGreaterEqual(len(catalog), 9)
        for entry in catalog:
            content, fname = generate_bzr_document(entry["kind"], self.company)
            self.assertIsNotNone(content, entry["kind"])
            self.assertGreater(len(content), 1000, entry["kind"])
            text = doc_text(content)
            self.assertNotIn("—", text, f"dash in {entry['kind']}")
            self.assertIn("Pekara Klas", text, entry["kind"])

    def test_unknown_kind_returns_none(self):
        content, fname = generate_bzr_document("nepostoj-x", self.company)
        self.assertIsNone(content)


class BzrDocumentsApiTest(APITestCase):
    def setUp(self):
        self.user = User.objects.create_superuser(
            "bzr", "bzr@example.com", "pass12345")
        self.client.force_authenticate(user=self.user)
        self.company = ClientCompany.objects.create(
            name="Firma BZR", tax_id=uuid.uuid4().hex[:9])

    def test_catalog_endpoint(self):
        r = self.client.get(
            "/api/partners/client-companies/bzr-documents-catalog/")
        self.assertEqual(r.status_code, 200)
        self.assertGreaterEqual(len(r.data), 9)

    def test_generate_endpoint_returns_docx(self):
        r = self.client.get(
            f"/api/partners/client-companies/{self.company.id}/bzr-document/",
            {"kind": "pravilnik"})
        self.assertEqual(r.status_code, 200)
        self.assertIn("wordprocessingml", r["Content-Type"])

    def test_generate_unknown_kind_400(self):
        r = self.client.get(
            f"/api/partners/client-companies/{self.company.id}/bzr-document/",
            {"kind": "xyz"})
        self.assertEqual(r.status_code, 400)
