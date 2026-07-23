import io
import uuid
from datetime import date

import docx
from django.contrib.auth import get_user_model
from django.core.files.base import ContentFile
from django.test import TestCase
from rest_framework.test import APIClient

from documents.models import DocumentCategory, DocumentTemplate
from partners.models import ClientCompany, Employee, JobRole
from processes.models import (
    ProcessBinding,
    ProcessRun,
    ProcessRunDocument,
    ProcessTemplate,
    ProcessType,
)

from .grading import grade_answers
from .models import TestAttempt, TestQuestion

User = get_user_model()


def make_company(**kwargs):
    defaults = {"name": "Test firma", "tax_id": uuid.uuid4().hex[:9]}
    defaults.update(kwargs)
    return ClientCompany.objects.create(**defaults)


def make_employee(company=None, **kwargs):
    defaults = {
        "client_company": company,
        "first_name": "Marko",
        "last_name": "Markovic",
    }
    defaults.update(kwargs)
    return Employee.objects.create(**defaults)


def make_question(order, correct_key="a", client_company=None):
    return TestQuestion.objects.create(
        client_company=client_company,
        text=f"Pitanje {order}?",
        choices=[
            {"key": "a", "text": "Tacno"},
            {"key": "b", "text": "Netacno"},
        ],
        correct_key=correct_key,
        order=order,
    )


def build_cell_docx() -> bytes:
    doc = docx.Document()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "label"
    table.cell(0, 1).text = ""
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


class GradeAnswersTest(TestCase):
    def setUp(self):
        self.employee = make_employee(make_company())

    def test_exact_75_percent_boundary_passes(self):
        questions = [make_question(i, "a") for i in range(1, 5)]
        answers = {str(q.id): "a" for q in questions[:3]}
        answers[str(questions[3].id)] = "b"
        score_pct, passed = grade_answers(self.employee, answers)
        self.assertEqual(score_pct, 75.0)
        self.assertTrue(passed)

    def test_below_75_percent_fails(self):
        questions = [make_question(i, "a") for i in range(1, 5)]
        answers = {str(q.id): "a" for q in questions[:2]}
        for q in questions[2:]:
            answers[str(q.id)] = "b"
        score_pct, passed = grade_answers(self.employee, answers)
        self.assertEqual(score_pct, 50.0)
        self.assertFalse(passed)

    def test_non_round_score_below_threshold_fails(self):
        questions = [make_question(i, "a") for i in range(1, 4)]
        answers = {str(questions[0].id): "a", str(
            questions[1].id): "b", str(questions[2].id): "b"}
        score_pct, passed = grade_answers(self.employee, answers)
        self.assertAlmostEqual(score_pct, 33.33)
        self.assertFalse(passed)

    def test_missing_answers_count_as_wrong(self):
        questions = [make_question(i, "a") for i in range(1, 5)]
        answers = {str(questions[0].id): "a"}
        score_pct, passed = grade_answers(self.employee, answers)
        self.assertEqual(score_pct, 25.0)
        self.assertFalse(passed)

    def test_global_and_client_scoped_questions_are_combined(self):
        make_question(1, "a", client_company=None)
        other_company = make_company(name="Druga firma")
        make_question(2, "a", client_company=other_company)
        own_question = make_question(3, "a", client_company=self.employee.client_company)
        answers = {str(own_question.id): "a"}
        score_pct, passed = grade_answers(self.employee, answers)
        self.assertEqual(score_pct, 50.0)
        self.assertFalse(passed)


class TestQuestionApiTest(TestCase):
    def setUp(self):
        self.company = make_company()
        self.question = make_question(1, "a")
        self.staff_user = User.objects.create_user(
            username="staffuser", password="x", is_staff=True,
        )
        self.plain_user = User.objects.create_user(
            username="plainuser", password="x", is_staff=False,
        )

    def test_correct_key_hidden_for_non_staff(self):
        client = APIClient()
        client.force_authenticate(user=self.plain_user)
        response = client.get("/api/testing/questions/")
        self.assertEqual(response.status_code, 200)
        row = response.data["results"][0]
        self.assertNotIn("correct_key", row)

    def test_correct_key_visible_for_staff(self):
        client = APIClient()
        client.force_authenticate(user=self.staff_user)
        response = client.get("/api/testing/questions/")
        self.assertEqual(response.status_code, 200)
        row = response.data["results"][0]
        self.assertEqual(row["correct_key"], "a")

    def test_client_company_filter_includes_global_and_scoped(self):
        scoped_company = make_company(name="Filter firma")
        make_question(2, "a", client_company=scoped_company)
        make_question(3, "a", client_company=make_company(name="Druga"))
        client = APIClient()
        client.force_authenticate(user=self.plain_user)
        response = client.get(
            f"/api/testing/questions/?client_company_id={scoped_company.id}"
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.data["results"]), 2)


class TestAttemptApiCompletionTest(TestCase):
    def setUp(self):
        self.user = User.objects.create_superuser(
            username="testingadmin", password="x", email="testingadmin@test.local",
        )
        self.category = DocumentCategory.objects.create(name="Test kategorija")
        self.doc_template = DocumentTemplate.objects.create(
            name="Obrazac 6 — evidencija o osposobljenosti za bezbedan rad",
            context_type=DocumentTemplate.CONTEXT_EMPLOYEE,
            category=self.category,
            generation_config={
                "mode": "DOCX_CELL_MAP",
                "cells": [
                    {"table": 0, "row": 0, "col": 1, "fieldKey": "employee.full_name"},
                ],
            },
        )
        self.process_type = ProcessType.objects.create(
            code="OSPOSOBLJAVANJE_BZR",
            name="Osposobljavanje BZR",
            subject_kind=ProcessType.SUBJECT_EMPLOYEE,
            default_period_months=36,
        )
        self.process_template = ProcessTemplate.objects.create(
            process_type=self.process_type,
            trigger=ProcessTemplate.TRIGGER_ON_COMPLETED,
            document_template=self.doc_template,
            generate_document=True,
        )
        self.company = make_company()
        self.role = JobRole.objects.create(
            client_company=self.company, name="Radno mesto",
        )
        self.role.obrazac6_template.save(
            "obrazac6_blank.docx", ContentFile(build_cell_docx()), save=True,
        )
        self.employee = make_employee(self.company, job_role=self.role)
        self.binding = ProcessBinding.objects.create(
            process_type=self.process_type,
            subject_kind=ProcessBinding.SUBJECT_EMPLOYEE,
            employee=self.employee,
            is_active=True,
            next_run_at=date.today(),
        )
        self.run = ProcessRun.objects.create(
            process_binding=self.binding,
            process_type=self.process_type,
            scheduled_for=date.today(),
            status=ProcessRun.STATUS_PENDING,
        )
        self.questions = [make_question(i, "a") for i in range(1, 5)]

    def _post_attempt(self, answers):
        client = APIClient()
        client.force_authenticate(user=self.user)
        return client.post(
            "/api/testing/attempts/",
            {"employee": self.employee.id, "answers": answers},
            format="json",
        )

    def test_passing_attempt_completes_run_and_generates_document(self):
        answers = {str(q.id): "a" for q in self.questions[:3]}
        answers[str(self.questions[3].id)] = "b"
        response = self._post_attempt(answers)
        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.data["score_pct"], 75.0)
        self.assertTrue(response.data["passed"])

        self.run.refresh_from_db()
        self.assertEqual(self.run.status, ProcessRun.STATUS_COMPLETED)
        self.assertTrue(
            ProcessRunDocument.objects.filter(process_run=self.run).exists()
        )

        attempt = TestAttempt.objects.get(id=response.data["id"])
        self.assertEqual(attempt.run_id, self.run.id)

    def test_failing_attempt_leaves_run_open(self):
        answers = {str(q.id): "b" for q in self.questions}
        response = self._post_attempt(answers)
        self.assertEqual(response.status_code, 201)
        self.assertFalse(response.data["passed"])

        self.run.refresh_from_db()
        self.assertEqual(self.run.status, ProcessRun.STATUS_PENDING)
        self.assertFalse(
            ProcessRunDocument.objects.filter(process_run=self.run).exists()
        )

        attempt = TestAttempt.objects.get(id=response.data["id"])
        self.assertIsNone(attempt.run)
