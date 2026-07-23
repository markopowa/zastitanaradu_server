import shutil
import tempfile
from pathlib import Path

import docx
from django.test import SimpleTestCase

from .tag_placement_migration import extract_tag_plan, has_tags, strip_tags_from_docx


def build_docx(build_fn) -> Path:
    doc = docx.Document()
    build_fn(doc)
    tmp_dir = Path(tempfile.mkdtemp(prefix="test_tag_migration_"))
    path = tmp_dir / "source.docx"
    doc.save(str(path))
    return path


class ExtractTagPlanTests(SimpleTestCase):
    def test_no_tags_is_empty_plan(self):
        path = build_docx(lambda doc: doc.add_paragraph("Bez tagova."))
        self.addCleanup(shutil.rmtree, path.parent, ignore_errors=True)
        document = docx.Document(str(path))
        self.assertFalse(has_tags(document))
        plan = extract_tag_plan(document)
        self.assertEqual(plan["occurrences"], [])
        self.assertEqual(plan["series_rows"], [])

    def test_paragraph_occurrences_in_reading_order_with_duplicates(self):
        def build(doc):
            doc.add_paragraph("Ime: {{ employee.full_name }}")
            doc.add_paragraph("Opet: {{ employee.full_name }} i {{ client.name }}")

        path = build_docx(build)
        self.addCleanup(shutil.rmtree, path.parent, ignore_errors=True)
        document = docx.Document(str(path))
        self.assertTrue(has_tags(document))
        plan = extract_tag_plan(document)
        self.assertEqual(
            plan["occurrences"],
            ["employee.full_name", "employee.full_name", "client.name"],
        )
        self.assertEqual(plan["series_rows"], [])

    def test_series_row_detected_and_excluded_from_occurrences(self):
        def build(doc):
            doc.add_paragraph("Firma: {{ client.name }}")
            table = doc.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "Redni broj"
            table.cell(0, 1).text = "Ime"
            table.cell(1, 0).text = "{{r.rbr}}"
            table.cell(1, 1).text = "{{r.employee_name}}"

        path = build_docx(build)
        self.addCleanup(shutil.rmtree, path.parent, ignore_errors=True)
        document = docx.Document(str(path))
        plan = extract_tag_plan(document)
        self.assertEqual(plan["occurrences"], ["client.name"])
        self.assertEqual(len(plan["series_rows"]), 1)
        series_row = plan["series_rows"][0]
        self.assertEqual(series_row["table"], 0)
        self.assertEqual(series_row["row"], 1)
        self.assertEqual(
            series_row["columns"], [(0, "r.rbr"), (1, "r.employee_name")],
        )

    def test_merged_cell_tag_counted_once_not_per_spanned_column(self):
        def build(doc):
            table = doc.add_table(rows=1, cols=4)
            table.cell(0, 2).merge(table.cell(0, 3))
            table.cell(0, 2).text = "{{ employee.full_name }}"

        path = build_docx(build)
        self.addCleanup(shutil.rmtree, path.parent, ignore_errors=True)
        document = docx.Document(str(path))
        # python-docx repeats the merged cell object for both spanned
        # columns; make sure that doesn't get double-counted.
        self.assertTrue(
            document.tables[0].rows[0].cells[2]._tc
            is document.tables[0].rows[0].cells[3]._tc
        )
        plan = extract_tag_plan(document)
        self.assertEqual(plan["occurrences"], ["employee.full_name"])

    def test_table_cell_tags_outside_series_row_counted_as_occurrences(self):
        def build(doc):
            table = doc.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "{{ employee.full_name }}"
            table.cell(0, 1).text = "{{ employee.full_name }}"

        path = build_docx(build)
        self.addCleanup(shutil.rmtree, path.parent, ignore_errors=True)
        document = docx.Document(str(path))
        plan = extract_tag_plan(document)
        self.assertEqual(
            plan["occurrences"], ["employee.full_name", "employee.full_name"],
        )
        self.assertEqual(plan["series_rows"], [])


class StripTagsFromDocxTests(SimpleTestCase):
    def test_strips_paragraph_tags_in_place(self):
        path = build_docx(
            lambda doc: doc.add_paragraph("Ime: {{ employee.full_name }} kraj."),
        )
        self.addCleanup(shutil.rmtree, path.parent, ignore_errors=True)

        changed = strip_tags_from_docx(path)
        self.assertTrue(changed)

        result = docx.Document(str(path))
        self.assertEqual(result.paragraphs[0].text, "Ime:  kraj.")
        self.assertNotIn("{{", result.paragraphs[0].text)

    def test_strips_table_cell_tags_recursively(self):
        def build(doc):
            table = doc.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "label"
            table.cell(0, 1).text = "{{ employee.full_name }}"

        path = build_docx(build)
        self.addCleanup(shutil.rmtree, path.parent, ignore_errors=True)

        changed = strip_tags_from_docx(path)
        self.assertTrue(changed)
        result = docx.Document(str(path))
        self.assertEqual(result.tables[0].cell(0, 1).text, "")

    def test_no_tags_is_noop_and_reports_unchanged(self):
        path = build_docx(lambda doc: doc.add_paragraph("Bez tagova."))
        self.addCleanup(shutil.rmtree, path.parent, ignore_errors=True)
        mtime_before = path.stat().st_mtime_ns

        changed = strip_tags_from_docx(path)
        self.assertFalse(changed)
        self.assertEqual(path.stat().st_mtime_ns, mtime_before)

    def test_preserves_surrounding_run_formatting(self):
        def build(doc):
            para = doc.add_paragraph()
            run = para.add_run("Pre: ")
            run.bold = True
            para.add_run("{{ employee.full_name }}")
            run2 = para.add_run(" posle.")
            run2.bold = True

        path = build_docx(build)
        self.addCleanup(shutil.rmtree, path.parent, ignore_errors=True)

        strip_tags_from_docx(path)
        result = docx.Document(str(path))
        para = result.paragraphs[0]
        self.assertEqual(para.text, "Pre:  posle.")
        self.assertTrue(para.runs[0].font.bold)
