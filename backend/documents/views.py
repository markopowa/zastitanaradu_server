import re
import docx
from pathlib import Path

from rest_framework import permissions, status, viewsets
from rest_framework.decorators import action
from rest_framework.response import Response

from .models import DocumentAIFormat, DocumentCategory, DocumentFile, DocumentTemplate
from .serializers import (
    DocumentAIFormatSerializer,
    DocumentCategorySerializer,
    DocumentFileSerializer,
    DocumentTemplateSerializer,
)


class DocumentCategoryViewSet(viewsets.ModelViewSet):
    queryset = DocumentCategory.objects.all().order_by("id")
    serializer_class = DocumentCategorySerializer
    permission_classes = [permissions.DjangoModelPermissions]


class DocumentFileViewSet(viewsets.ModelViewSet):
    queryset = DocumentFile.objects.select_related("category", "uploaded_by").all().order_by(
        "-uploaded_at"
    )
    serializer_class = DocumentFileSerializer
    permission_classes = [permissions.DjangoModelPermissions]

    def perform_create(self, serializer):
        serializer.save(uploaded_by=self.request.user)


class DocumentTemplateViewSet(viewsets.ModelViewSet):
    queryset = DocumentTemplate.objects.select_related(
        "category").all().order_by("id")
    serializer_class = DocumentTemplateSerializer
    permission_classes = [permissions.DjangoModelPermissions]

    @action(detail=False, methods=["post"], url_path="from-document")
    def from_document(self, request, *args, **kwargs):
        document_file_id = request.data.get("document_file_id")
        context_type = request.data.get("context_type")
        if not document_file_id or not context_type:
            return Response(
                {"detail": "document_file_id and context_type are required."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        try:
            document_file = DocumentFile.objects.get(pk=document_file_id)
        except DocumentFile.DoesNotExist:
            return Response(
                {"detail": "Document file not found."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        file_field = getattr(document_file, "file", None)
        if not file_field:
            return Response(
                {"detail": "Document file has no associated file."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        path = Path(file_field.path)
        suffix = path.suffix.lower()
        body = ""

        underline_pattern = re.compile(r"_{3,}")

        if suffix == ".docx":
            doc = docx.Document(path)

            lines: list[str] = []
            counter = 1

            for para in doc.paragraphs:
                text = para.text or ""

                def repl(match: re.Match[str]) -> str:
                    nonlocal counter
                    token = f"{{{{ field_{counter} }}}}"
                    counter += 1
                    return token

                converted = underline_pattern.sub(repl, text)
                lines.append(converted)

            body = "\n".join(lines)
        elif suffix in {".txt", ".jinja", ".jinja2"}:
            with open(path, "r", encoding="utf-8", errors="ignore") as f:
                raw = f.read()
            counter = 1

            def repl(match: re.Match[str]) -> str:
                nonlocal counter
                token = f"{{{{ field_{counter} }}}}"
                counter += 1
                return token

            body = underline_pattern.sub(repl, raw)
        else:
            return Response(
                {
                    "detail": "Unsupported file type for template generation.",
                    "extension": suffix,
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        payload = {
            "name": request.data.get("name") or document_file.title,
            "description": request.data.get("description") or "",
            "category_id": request.data.get("category_id"),
            "context_type": context_type,
            "template_body": body or "",
        }

        serializer = self.get_serializer(data=payload)
        serializer.is_valid(raise_exception=True)
        instance: DocumentTemplate = serializer.save(
            source_document_file_id=document_file.id)

        file_field = getattr(document_file, "file", None)
        if file_field:
            # Save a copy of the original file as the template file
            instance.template_file.save(
                Path(file_field.name).name, file_field.file, save=True)

        out = self.get_serializer(instance)
        headers = self.get_success_headers(out.data)
        return Response(out.data, status=status.HTTP_201_CREATED, headers=headers)

    @action(detail=False, methods=["post"], url_path="from-file")
    def from_file(self, request, *args, **kwargs):
        uploaded_file = request.FILES.get("file")
        context_type = request.data.get("context_type")
        if not uploaded_file or not context_type:
            return Response(
                {"detail": "file and context_type are required."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        name = request.data.get("name") or getattr(
            uploaded_file, "name", "") or "Novi šablon"
        description = request.data.get("description") or ""
        category_id = request.data.get("category_id")

        suffix = Path(uploaded_file.name).suffix.lower()

        underline_pattern = re.compile(r"_{3,}")
        body = ""

        if suffix == ".docx":
            # python-docx can work directly with the uploaded file object
            doc = docx.Document(uploaded_file)
            lines: list[str] = []
            counter = 1

            for para in doc.paragraphs:
                text = para.text or ""

                def repl(match: re.Match[str]) -> str:
                    nonlocal counter
                    token = f"{{{{ field_{counter} }}}}"
                    counter += 1
                    return token

                converted = underline_pattern.sub(repl, text)
                lines.append(converted)

            body = "\n".join(lines)
        elif suffix in {".txt", ".jinja", ".jinja2"}:
            raw_bytes = uploaded_file.read()
            raw = raw_bytes.decode("utf-8", errors="ignore")
            counter = 1

            def repl(match: re.Match[str]) -> str:
                nonlocal counter
                token = f"{{{{ field_{counter} }}}}"
                counter += 1
                return token

            body = underline_pattern.sub(repl, raw)
        else:
            # For other file types we still accept the upload but do not attempt to build template_body
            body = request.data.get("template_body") or ""

        payload = {
            "name": name,
            "description": description,
            "category_id": category_id,
            "context_type": context_type,
            "template_body": body or "",
        }

        serializer = self.get_serializer(data=payload)
        serializer.is_valid(raise_exception=True)
        instance: DocumentTemplate = serializer.save()
        # Reset file pointer before saving to FileField storage
        try:
            uploaded_file.seek(0)
        except (AttributeError, OSError):
            pass
        instance.template_file.save(
            uploaded_file.name, uploaded_file, save=True)
        out = self.get_serializer(instance)
        return Response(out.data, status=status.HTTP_201_CREATED)


class DocumentAIFormatViewSet(viewsets.ReadOnlyModelViewSet):
    queryset = DocumentAIFormat.objects.filter(is_active=True).order_by("id")
    serializer_class = DocumentAIFormatSerializer
    permission_classes = [permissions.IsAuthenticated]
