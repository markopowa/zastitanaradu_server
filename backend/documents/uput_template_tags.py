from docx.oxml.ns import qn

from .utils import _copy_run_format, _path_str

UPUT_BODY_PRETHODNI = (
    "Upućuje se na PRETHODNI pregled {{ employee.first_name }} "
    "{{ employee.father_name }} {{ employee.last_name }}, rođen(a) "
    "{{ date_of_birth }} u {{ employee.place_of_birth }}, JMBG "
    "{{ employee.national_id }}, po zanimanju {{ employee.occupation }}, "
    "koji(a) treba da radi na radnom mestu "
    "{{ employee.high_risk_position_name }} kod poslodavca {{ client.name }}, "
    "radi ocene ispunjenosti posebnih zdravstvenih sposobnosti za obavljanje "
    "poslova na tom radnom mestu – koje je Aktom o proceni rizika "
    "({{ client.risk_assessment_act_name }}, "
    "{{ client.risk_assessment_act_date }}) utvrđeno kao radno mesto sa "
    "povećanim rizikom."
)

UPUT_BODY_PERIODICNI = (
    "Upućuje se na PERIODIČNI/KONTROLNI pregled {{ employee.first_name }} "
    "{{ employee.father_name }} {{ employee.last_name }}, rođen(a) "
    "{{ date_of_birth }} u {{ employee.place_of_birth }}, JMBG "
    "{{ employee.national_id }}, po zanimanju {{ employee.occupation }}, "
    "koji(a) treba da radi na radnom mestu "
    "{{ employee.high_risk_position_name }} kod poslodavca {{ client.name }}, "
    "radi ocene ispunjenosti posebnih zdravstvenih sposobnosti za obavljanje "
    "poslova na tom radnom mestu – koje je Aktom o proceni rizika "
    "({{ client.risk_assessment_act_name }}, "
    "{{ client.risk_assessment_act_date }}) utvrđeno kao radno mesto sa "
    "povećanim rizikom."
)

UPUT_LAST_EXAM = (
    "Pri prethodnom/periodičnom pregledu obavljenom {{ last_exam_date }} "
    "u zdravstvenoj ustanovi _______________________________ - službi medicine "
    "rada, utvrđeno je: _______________________________________________"
)


def _set_paragraph_text(para, text: str) -> None:
    base_rpr = None
    if para.runs:
        base_rpr = para.runs[0]._element.find(qn("w:rPr"))
    para.clear()
    run = para.add_run(text)
    _copy_run_format(run, base_rpr)


def restore_uput_tags(document, *, periodic: bool) -> bool:
    changed = False
    for para in document.paragraphs:
        text = (para.text or "").strip()
        if text.startswith("Poslodavac:"):
            _set_paragraph_text(para, "Poslodavac: {{ client.name }}")
            changed = True
        elif text.startswith("Matični broj iz jedinstvenog registra:"):
            _set_paragraph_text(
                para,
                "Matični broj iz jedinstvenog registra: "
                "{{ client.registration_number }}",
            )
            changed = True
        elif text.startswith("Adresa:"):
            _set_paragraph_text(para, "Adresa: {{ client.address }}")
            changed = True
        elif text.startswith("Šifra delatnosti:"):
            _set_paragraph_text(
                para, "Šifra delatnosti: {{ client.activity_code }}")
            changed = True
        elif text.startswith("Datum:"):
            _set_paragraph_text(para, "Datum: {{ scheduled_for }}")
            changed = True
        elif text.startswith("Broj uputa:"):
            _set_paragraph_text(para, "Broj uputa: {{ instruction_number }}")
            changed = True
        elif text.startswith("Upućuje se na PRETHODNI pregled"):
            _set_paragraph_text(para, UPUT_BODY_PRETHODNI)
            changed = True
        elif text.startswith("Upućuje se na PERIODIČNI/KONTROLNI pregled"):
            _set_paragraph_text(para, UPUT_BODY_PERIODICNI)
            changed = True
        elif periodic and text.startswith(
                "Pri prethodnom/periodičnom pregledu obavljenom"):
            _set_paragraph_text(para, UPUT_LAST_EXAM)
            changed = True
    return changed


def restore_uput_tags_file(path, *, periodic: bool) -> bool:
    import docx

    document = docx.Document(_path_str(path))
    changed = restore_uput_tags(document, periodic=periodic)
    if changed:
        document.save(_path_str(path))
    return changed
