import io

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

LABELS = {
    "client.name": "Naziv firme",
    "client.tax_id": "PIB",
    "client.registration_number": "Matični broj",
    "client.address": "Adresa firme",
    "client.activity_code": "Šifra delatnosti",
    "client.phone": "Telefon firme",
    "client.email": "Email firme",
    "client.website": "Web sajt firme",
    "client.risk_assessment_act_name": "Naziv Akta o proceni rizika",
    "client.risk_assessment_act_date": "Datum donošenja Akta o proceni rizika",
    "employee.first_name": "Ime zaposlenog",
    "employee.last_name": "Prezime zaposlenog",
    "employee.father_name": "Ime oca",
    "employee.national_id": "JMBG",
    "employee.date_of_birth": "Datum rođenja",
    "employee.place_of_birth": "Mesto rođenja",
    "employee.occupation": "Zanimanje",
    "employee.org_unit": "Organizaciona jedinica",
    "employee.position": "Radno mesto",
    "employee.email": "Email zaposlenog",
    "employee.high_risk_position_name": "Radno mesto sa povećanim rizikom",
    "equipment.name": "Naziv opreme",
    "equipment.category": "Kategorija opreme",
    "equipment.inventory_number": "Inventarski broj",
    "equipment.location": "Lokacija opreme",
    "scheduled_for": "Datum",
    "performed_at": "Datum izvođenja",
    "valid_until": "Važi do",
    "process_type_name": "Vrsta obaveze",
    "instruction_number": "Broj uputa",
    "last_exam_date": "Datum prethodnog pregleda",
    "date_of_birth": "Datum rođenja",
    "year_of_birth": "Godina rođenja",
    "training_type_name": "Vrsta obuke",
}

FIXED_TEXT_KEY = "__fixed_text__"


def label_for(key):
    if key in LABELS:
        return LABELS[key]
    return key.split(".")[-1].replace("_", " ").capitalize()


def _order(key):
    if key.startswith("client."):
        return (0, key)
    if key.startswith("employee."):
        return (1, key)
    if key.startswith("equipment."):
        return (2, key)
    return (3, key)


def _set(cell, text, bold=False, size=11):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(text if text is not None else "")
    r.bold = bold
    r.font.size = Pt(size)


def build_form_template(title, intro, field_keys):
    keys = [k for k in field_keys if k and k != FIXED_TEXT_KEY]
    seen = set()
    ordered = []
    for k in sorted(keys, key=_order):
        if k not in seen:
            seen.add(k)
            ordered.append(k)

    doc = docx.Document()
    section = doc.sections[0]
    section.left_margin = section.right_margin = Cm(2.0)
    section.top_margin = section.bottom_margin = Cm(2.0)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(14)

    if intro:
        doc.add_paragraph()
        ip = doc.add_paragraph()
        ir = ip.add_run(intro)
        ir.font.size = Pt(11)

    doc.add_paragraph()

    table = doc.add_table(rows=len(ordered), cols=2)
    table.style = "Table Grid"
    table.columns[0].width = Cm(6.5)
    table.columns[1].width = Cm(9.5)
    for i, key in enumerate(ordered):
        row = table.rows[i]
        for c in row.cells:
            c.width = Cm(6.5) if c is row.cells[0] else Cm(9.5)
        _set(row.cells[0], label_for(key), bold=True, size=10)
        _set(row.cells[1], f"{{{{{key}}}}}", size=11)

    doc.add_paragraph()
    doc.add_paragraph()
    sig = doc.add_table(rows=1, cols=2)
    sig.cell(0, 0).text = (
        "_____________________________\n"
        "Lice za bezbednost i zdravlje na radu"
    )
    sig.cell(0, 1).text = "_____________________________\nPoslodavac"

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


TEMPLATE_INTROS = {
    "uput_lekarski": (
        "Poslodavac upućuje zaposlenog na periodični lekarski pregled u "
        "skladu sa Zakonom o bezbednosti i zdravlju na radu i Pravilnikom o "
        "prethodnim i periodičnim lekarskim pregledima zaposlenih na radnim "
        "mestima sa povećanim rizikom. Podaci o zaposlenom i radnom mestu:"
    ),
    "uput_prethodni": (
        "Poslodavac upućuje zaposlenog na prethodni lekarski pregled u skladu "
        "sa Zakonom o bezbednosti i zdravlju na radu i Pravilnikom o "
        "prethodnim i periodičnim lekarskim pregledima zaposlenih na radnim "
        "mestima sa povećanim rizikom. Podaci o zaposlenom i radnom mestu:"
    ),
    "obrazac6": (
        "Evidencija o osposobljenosti zaposlenog za bezbedan i zdrav rad "
        "(Obrazac 6), u skladu sa propisima o vođenju evidencija u oblasti "
        "bezbednosti i zdravlja na radu. Podaci:"
    ),
    "lzo_revers": (
        "Karton o zaduženju sredstvima i opremom za ličnu zaštitu na radu "
        "(LZO). Podaci o zaposlenom i zaduženju:"
    ),
    "potvrda_clan5": (
        "Potvrda o osposobljenosti zaposlenog za bezbedan i zdrav rad. "
        "Podaci:"
    ),
}
