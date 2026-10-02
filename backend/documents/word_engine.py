import copy
import io
import re

import docx
from docx.oxml.ns import qn

TAG_RE = re.compile(r"\{\{\s*([^{}#/]+?)\s*\}\}")
EACH_RE = re.compile(r"^\{\{\s*#each\s+([\w.]+)\s+as\s+(\w+)\s*\}\}$")
IF_RE = re.compile(r"^\{\{\s*#if\s+(not\s+)?([\w.]+)\s*\}\}$")
END_RE = re.compile(r"^\{\{\s*/(each|if)\s*\}\}$")

W_P = qn("w:p")
W_R = qn("w:r")
W_T = qn("w:t")
W_TR = qn("w:tr")
W_TBL = qn("w:tbl")
W_TC = qn("w:tc")
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"


class TemplateError(Exception):
    pass


def lookup(path, context):
    current = context
    for part in path.split("."):
        if isinstance(current, dict):
            current = current.get(part)
        else:
            return None
        if current is None:
            return None
    return current


def _text_value(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Da" if value else "Ne"
    return str(value)


W_TXBX = qn("w:txbxContent")


def _own_runs(p):
    for r in p.iter(W_R):
        parent = r.getparent()
        nested = False
        while parent is not None and parent is not p:
            if parent.tag == W_TXBX:
                nested = True
                break
            parent = parent.getparent()
        if not nested:
            yield r


def _run_texts(p):
    nodes = []
    for r in _own_runs(p):
        nodes.extend(r.findall(W_T))
    return nodes


def _paragraph_text(p):
    return "".join(t.text or "" for t in _run_texts(p))


def _replace_in_paragraph(p, context):
    nodes = _run_texts(p)
    if not nodes:
        return
    full = "".join(n.text or "" for n in nodes)
    if "{{" not in full:
        return
    matches = list(TAG_RE.finditer(full))
    if not matches:
        return
    spans = []
    pos = 0
    for n in nodes:
        length = len(n.text or "")
        spans.append((pos, pos + length))
        pos += length
    for m in reversed(matches):
        start, end = m.span()
        value = _text_value(lookup(m.group(1).strip(), context))
        first = True
        for node, (s, e) in zip(nodes, spans):
            if e <= start or s >= end:
                continue
            text = node.text or ""
            local_start = max(start, s) - s
            local_end = min(end, e) - s
            if first:
                node.text = text[:local_start] + value + text[local_end:]
                first = False
            else:
                node.text = text[:local_start] + text[local_end:]
            node.set(XML_SPACE, "preserve")
        pos = 0
        spans = []
        for n in nodes:
            length = len(n.text or "")
            spans.append((pos, pos + length))
            pos += length
    _expand_line_breaks(nodes)


def _expand_line_breaks(nodes):
    for node in nodes:
        text = node.text or ""
        if "\n" not in text:
            continue
        run = node.getparent()
        parts = text.split("\n")
        node.text = parts[0]
        anchor = node
        for part in parts[1:]:
            br = run.makeelement(qn("w:br"), {})
            anchor.addnext(br)
            t = run.makeelement(W_T, {})
            t.text = part
            t.set(XML_SPACE, "preserve")
            br.addnext(t)
            anchor = t


def _row_list_binding(tr, context):
    for p in tr.iter(W_P):
        for m in TAG_RE.finditer(_paragraph_text(p)):
            parts = m.group(1).strip().split(".")
            current = context
            for i, part in enumerate(parts):
                current = current.get(part) if isinstance(current, dict) else None
                if isinstance(current, list):
                    return ".".join(parts[: i + 1]), current
                if current is None:
                    break
    return None, None


def _bind(context, path, item):
    bound = dict(context)
    parts = path.split(".")
    if len(parts) == 1:
        bound[parts[0]] = item
        return bound
    head = dict(bound.get(parts[0]) or {})
    bound[parts[0]] = head
    node = head
    for part in parts[1:-1]:
        child = dict(node.get(part) or {})
        node[part] = child
        node = child
    node[parts[-1]] = item
    return bound


def _render_table(tbl, context):
    for tr in list(tbl.findall(W_TR)):
        path, items = _row_list_binding(tr, context)
        if path is None:
            _render_cells(tr, context)
            continue
        anchor = tr
        for index, item in enumerate(items, start=1):
            row_item = dict(item) if isinstance(item, dict) else {"value": item}
            row_item.setdefault("rbr", f"{index}.")
            clone = copy.deepcopy(tr)
            _render_cells(clone, _bind(context, path, row_item))
            anchor.addnext(clone)
            anchor = clone
        tr.getparent().remove(tr)


def _render_cells(tr, context):
    for tc in tr.findall(W_TC):
        _render_container(tc, context)


def _render_container(container, context):
    _render_children(list(container), context)


def _render_children(children, context):
    index = 0
    while index < len(children):
        child = children[index]
        marker = _paragraph_text(child).strip() if child.tag == W_P else ""
        each = EACH_RE.match(marker)
        cond = IF_RE.match(marker)
        if each or cond:
            end_index = _matching_end(children, index)
            body = children[index + 1:end_index]
            parent = child.getparent()
            anchor = children[end_index]
            if each:
                for item in lookup(each.group(1), context) or []:
                    clones = [copy.deepcopy(element) for element in body]
                    for clone in clones:
                        anchor.addprevious(clone)
                    _render_children(clones, {**context, each.group(2): item})
            else:
                truthy = bool(lookup(cond.group(2), context))
                if cond.group(1):
                    truthy = not truthy
                if truthy:
                    clones = [copy.deepcopy(element) for element in body]
                    for clone in clones:
                        anchor.addprevious(clone)
                    _render_children(clones, context)
            for element in [child, *body, anchor]:
                parent.remove(element)
            index = end_index + 1
            continue
        _render_element(child, context)
        index += 1


def _matching_end(children, start):
    depth = 0
    for i in range(start, len(children)):
        child = children[i]
        if child.tag != W_P:
            continue
        text = _paragraph_text(child).strip()
        if EACH_RE.match(text) or IF_RE.match(text):
            depth += 1
        elif END_RE.match(text):
            depth -= 1
            if depth == 0:
                return i
    raise TemplateError(f"Blok bez zatvaranja: {_paragraph_text(children[start])}")


def _render_element(element, context):
    if element.tag == W_P:
        _replace_in_paragraph(element, context)
        for nested in element.iter(W_TXBX):
            _render_container(nested, context)
    elif element.tag == W_TBL:
        _render_table(element, context)
    else:
        for p in element.iter(W_P):
            _replace_in_paragraph(p, context)


def _header_footer_parts(document):
    seen = set()
    for section in document.sections:
        for part in (
            section.header, section.footer,
            section.first_page_header, section.first_page_footer,
            section.even_page_header, section.even_page_footer,
        ):
            try:
                element = part._element
            except Exception:
                continue
            if id(element) in seen:
                continue
            seen.add(id(element))
            yield element


def render_document(document, context):
    _render_container(document.element.body, context)
    for element in _header_footer_parts(document):
        _render_container(element, context)
    return document


def render_docx(source, context) -> bytes:
    if isinstance(source, (bytes, bytearray)):
        document = docx.Document(io.BytesIO(source))
    else:
        document = docx.Document(source)
    render_document(document, context)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def template_tags(source) -> list[str]:
    if isinstance(source, (bytes, bytearray)):
        document = docx.Document(io.BytesIO(source))
    else:
        document = docx.Document(source)
    elements = [document.element.body, *_header_footer_parts(document)]
    tags = []
    for element in elements:
        for p in element.iter(W_P):
            for m in TAG_RE.finditer(_paragraph_text(p)):
                key = m.group(1).strip()
                if key not in tags:
                    tags.append(key)
    return tags
