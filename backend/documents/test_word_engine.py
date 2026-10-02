import io

import docx
from django.test import SimpleTestCase

from documents.word_engine import render_docx, template_tags


def _build(build):
    document = docx.Document()
    build(document)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _open(content):
    return docx.Document(io.BytesIO(content))


def _texts(document):
    return [p.text for p in document.paragraphs]


class WordEngineTest(SimpleTestCase):
    def test_tag_split_across_runs_keeps_first_run_format(self):
        def build(document):
            p = document.add_paragraph("Firma: ")
            bold = p.add_run("{{ cli")
            bold.bold = True
            p.add_run("ent.name }} kraj")

        result = _open(render_docx(_build(build), {"client": {"name": "MAK"}}))
        paragraph = result.paragraphs[0]
        self.assertEqual(paragraph.text, "Firma: MAK kraj")
        self.assertTrue(paragraph.runs[1].bold)

    def test_missing_value_renders_empty(self):
        content = _build(lambda d: d.add_paragraph("A {{ x.y }} B"))
        self.assertEqual(_texts(_open(render_docx(content, {})))[0], "A  B")

    def test_header_and_footer_are_filled(self):
        def build(document):
            document.sections[0].header.paragraphs[0].text = "{{ client.name }}"
            document.sections[0].footer.paragraphs[0].text = "PIB {{ client.tax_id }}"

        result = _open(render_docx(
            _build(build), {"client": {"name": "MAK", "tax_id": "123"}}))
        self.assertEqual(result.sections[0].header.paragraphs[0].text, "MAK")
        self.assertEqual(result.sections[0].footer.paragraphs[0].text, "PIB 123")

    def test_table_row_repeats_per_item_with_numbering(self):
        def build(document):
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "R.br."
            table.cell(0, 1).text = "Ime"
            table.cell(1, 0).text = "{{ exams.rbr }}"
            table.cell(1, 1).text = "{{ exams.name }}"

        context = {"exams": [{"name": "Ana"}, {"name": "Petar"}]}
        table = _open(render_docx(_build(build), context)).tables[0]
        self.assertEqual(len(table.rows), 3)
        self.assertEqual(table.cell(1, 0).text, "1.")
        self.assertEqual(table.cell(2, 1).text, "Petar")

    def test_empty_list_removes_template_row(self):
        def build(document):
            table = document.add_table(rows=2, cols=1)
            table.cell(0, 0).text = "Ime"
            table.cell(1, 0).text = "{{ exams.name }}"

        table = _open(render_docx(_build(build), {"exams": []})).tables[0]
        self.assertEqual(len(table.rows), 1)

    def test_each_block_with_nested_rows(self):
        def build(document):
            document.add_paragraph("{{#each roles as role}}")
            document.add_paragraph("Radno mesto: {{ role.name }}")
            table = document.add_table(rows=2, cols=1)
            table.cell(0, 0).text = "Opasnost"
            table.cell(1, 0).text = "{{ role.hazards.label }}"
            document.add_paragraph("{{/each}}")
            document.add_paragraph("Kraj")

        context = {"roles": [
            {"name": "Magacioner", "hazards": [{"label": "Pad"}, {"label": "Teret"}]},
            {"name": "Vozač", "hazards": [{"label": "Saobraćaj"}]},
        ]}
        result = _open(render_docx(_build(build), context))
        texts = [t for t in _texts(result) if t]
        self.assertEqual(texts, ["Radno mesto: Magacioner", "Radno mesto: Vozač", "Kraj"])
        self.assertEqual([len(t.rows) for t in result.tables], [3, 2])
        self.assertEqual(result.tables[1].cell(1, 0).text, "Saobraćaj")

    def test_nested_blocks(self):
        def build(document):
            document.add_paragraph("{{#each roles as role}}")
            document.add_paragraph("{{ role.name }}")
            document.add_paragraph("{{#if role.high}}")
            document.add_paragraph("Povećan rizik")
            document.add_paragraph("{{/if}}")
            document.add_paragraph("{{#each role.hazards as hazard}}")
            document.add_paragraph("- {{ hazard.label }}")
            document.add_paragraph("{{/each}}")
            document.add_paragraph("{{/each}}")

        context = {"roles": [
            {"name": "A", "high": True, "hazards": [{"label": "x"}, {"label": "y"}]},
            {"name": "B", "high": False, "hazards": []},
        ]}
        result = _open(render_docx(_build(build), context))
        self.assertEqual(
            [t for t in _texts(result) if t],
            ["A", "Povećan rizik", "- x", "- y", "B"],
        )

    def test_if_block(self):
        def build(document):
            document.add_paragraph("{{#if client.high}}")
            document.add_paragraph("Povećan rizik")
            document.add_paragraph("{{/if}}")
            document.add_paragraph("{{#if not client.high}}")
            document.add_paragraph("Nema")
            document.add_paragraph("{{/if}}")

        result = _open(render_docx(_build(build), {"client": {"high": False}}))
        self.assertEqual([t for t in _texts(result) if t], ["Nema"])

    def test_multiline_value_becomes_line_breaks(self):
        content = _build(lambda d: d.add_paragraph("{{ text }}"))
        result = _open(render_docx(content, {"text": "a\nb"}))
        self.assertEqual(result.paragraphs[0].text, "a\nb")

    def test_template_tags_lists_unique_keys(self):
        def build(document):
            document.add_paragraph("{{ a }} {{ b.c }} {{ a }}")
            document.sections[0].header.paragraphs[0].text = "{{ d }}"

        self.assertEqual(template_tags(_build(build)), ["a", "b.c", "d"])
