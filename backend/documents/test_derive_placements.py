import shutil
import tempfile
from pathlib import Path

import docx
import fitz
from django.conf import settings
from django.test import TestCase

from .models import DocumentCategory, TemplateFieldDefinition
from .tag_placement_migration import derive_placements, strip_tags_from_docx


def build_docx(build_fn) -> Path:
    doc = docx.Document()
    build_fn(doc)
    tmp_dir = Path(tempfile.mkdtemp(prefix="test_derive_placements_"))
    path = tmp_dir / "source.docx"
    doc.save(str(path))
    return path


def require_libreoffice(test_case) -> None:
    if not shutil.which(settings.LIBREOFFICE_BIN):
        test_case.skipTest("LibreOffice not available on this host")


class DerivePlacementsTests(TestCase):
    def setUp(self):
        TemplateFieldDefinition.objects.create(
            key="employee.full_name",
            label="Puno ime",
            category=TemplateFieldDefinition.CATEGORY_EMPLOYEE,
            order=1,
        )
        TemplateFieldDefinition.objects.create(
            key="client.name",
            label="Naziv firme",
            category=TemplateFieldDefinition.CATEGORY_CLIENT_COMPANY,
            order=2,
        )
        self.category = DocumentCategory.objects.create(name="Test kategorija")

    def _locate_text_pct(self, pdf_bytes: bytes, needle: str):
        doc = fitz.open(stream=pdf_bytes, filetype="pdf")
        try:
            page = doc[0]
            rects = page.search_for(needle)
            self.assertTrue(rects, f"expected to find {needle!r} in rendered pdf")
            rect = rects[0]
            page_rect = page.rect
            return (
                rect.x0 / page_rect.width * 100,
                rect.y0 / page_rect.height * 100,
            )
        finally:
            doc.close()

    def test_two_tags_and_a_series_row_derive_within_tolerance(self):
        require_libreoffice(self)

        def build(doc):
            doc.add_paragraph("Ime: {{ employee.full_name }}")
            doc.add_paragraph("Firma: {{ client.name }}")
            table = doc.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "Redni broj"
            table.cell(0, 1).text = "Ime"
            table.cell(1, 0).text = "{{r.rbr}}"
            table.cell(1, 1).text = "{{r.employee_name}}"

        path = build_docx(build)
        self.addCleanup(shutil.rmtree, path.parent, ignore_errors=True)

        placements, warnings = derive_placements(
            path, series_source_map={0: "high_risk_employees_for_company"},
        )
        self.assertEqual(warnings, [])

        field_placements = [p for p in placements if "seriesSource" not in p]
        series_placements = [p for p in placements if "seriesSource" in p]
        self.assertEqual(len(field_placements), 2)
        self.assertEqual(len(series_placements), 1)

        by_key = {p["fieldKey"]: p for p in field_placements}
        self.assertIn("employee.full_name", by_key)
        self.assertIn("client.name", by_key)
        for ph in field_placements:
            self.assertGreaterEqual(ph["xPct"], 0)
            self.assertGreaterEqual(ph["yPct"], 0)
            self.assertGreater(ph["widthPct"], 0)
            self.assertGreater(ph["heightPct"], 0)
            self.assertGreaterEqual(ph["fontSize"], 6)
            self.assertLessEqual(ph["fontSize"], 48)

        series = series_placements[0]
        self.assertEqual(series["seriesSource"], "high_risk_employees_for_company")
        self.assertEqual(len(series["columns"]), 2)
        fields = {c["field"] for c in series["columns"]}
        self.assertEqual(fields, {"r.rbr", "r.employee_name"})
        # rbr column should render to the left of the employee_name column.
        by_field = {c["field"]: c for c in series["columns"]}
        self.assertLess(by_field["r.rbr"]["xPct"], by_field["r.employee_name"]["xPct"])

        # Re-derive after stripping to confirm idempotency: no tags left.
        changed = strip_tags_from_docx(path)
        self.assertTrue(changed)
        placements_after, _warnings_after = derive_placements(path)
        self.assertEqual(placements_after, [])

    def test_derived_positions_match_rendered_pdf_within_tolerance(self):
        require_libreoffice(self)

        def build(doc):
            doc.add_paragraph("Zaposleni: {{ employee.full_name }}")

        path = build_docx(build)
        self.addCleanup(shutil.rmtree, path.parent, ignore_errors=True)

        placements, warnings = derive_placements(path)
        self.assertEqual(warnings, [])
        self.assertEqual(len(placements), 1)
        ph = placements[0]

        from .utils import fill_pdf_at_coordinates

        context = {"employee": {"full_name": "Marko Markovic"}}
        result = fill_pdf_at_coordinates(path, [ph], context)
        x_pct, y_pct = self._locate_text_pct(result, "Marko Markovic")
        self.assertAlmostEqual(x_pct, ph["xPct"], delta=2.0)
        self.assertAlmostEqual(y_pct, ph["yPct"], delta=2.0)
