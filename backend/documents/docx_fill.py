import copy
import io
import re

import docx
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

TAG_RE = re.compile(r"\{\{\s*([^{}]+?)\s*\}\}")
RPR = qn("w:rPr")


def _resolve(key, context):
    cur = context
    for part in key.split("."):
        if isinstance(cur, dict):
            cur = cur.get(part)
        else:
            cur = None
        if cur is None:
            return ""
    return "" if cur is None else str(cur)


def iter_all_paragraphs(document):
    seen = set()
    for p in document.element.body.iter(qn("w:p")):
        if id(p) in seen:
            continue
        seen.add(id(p))
        yield Paragraph(p, document)
    for section in document.sections:
        for hf in (section.header, section.footer,
                   section.first_page_header, section.first_page_footer,
                   section.even_page_header, section.even_page_footer):
            try:
                el = hf._element
            except Exception:
                continue
            for p in el.iter(qn("w:p")):
                if id(p) in seen:
                    continue
                seen.add(id(p))
                yield Paragraph(p, hf)


def _fill_paragraph(para, context):
    text = para.text or ""
    if "{{" not in text:
        return
    new_text = TAG_RE.sub(lambda m: _resolve(m.group(1), context), text)
    if new_text == text:
        return
    base_rpr = para.runs[0]._element.find(RPR) if para.runs else None
    for run in list(para.runs):
        run._element.getparent().remove(run._element)
    lines = new_text.split("\n")
    run = para.add_run(lines[0])
    if base_rpr is not None:
        run._element.insert(0, copy.deepcopy(base_rpr))
    for line in lines[1:]:
        r = para.add_run()
        r.add_break()
        r2 = para.add_run(line)
        if base_rpr is not None:
            r2._element.insert(0, copy.deepcopy(base_rpr))


def fill_docx_placeholders(src_path, context) -> bytes:
    document = docx.Document(str(src_path))
    for para in iter_all_paragraphs(document):
        _fill_paragraph(para, context)
    buf = io.BytesIO()
    document.save(buf)
    buf.seek(0)
    return buf.read()
