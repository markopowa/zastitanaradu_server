import copy
import io
import logging
import re

import docx
from docx.table import _Row
from jinja2 import Template

from django.conf import settings
from django.contrib.auth import get_user_model
from django.core.files.base import ContentFile

from core.email_context import company_email_test_mode
from core.email_sender import get_email_sender
from documents.models import DocumentCategory, DocumentFile
from documents.utils import _resolve_field_value, generate_visual_pdf

from .date_format import format_date_display
from .models import (
    ProcessBinding,
    ProcessRun,
    ProcessRunDocument,
    ProcessTemplate,
    get_next_code,
)
from .result_data import (
    normalize_client_snapshot_tax_id,
    normalize_employee_snapshot_national_id,
)

logger = logging.getLogger(__name__)
User = get_user_model()


def _risk_assessment_act_name(client) -> str:
    f = getattr(client, "risk_assessment_act_file", None)
    if not f:
        return ""
    import os as _os

    base = _os.path.basename(getattr(f, "name", "") or "")
    name, _ = _os.path.splitext(base)
    return name


def _resolve_client_contacts(client) -> dict:
    try:
        persons = list(client.contact_persons.all())
    except Exception:
        persons = []

    directors = [p for p in persons if p.role == "DIRECTOR"]
    director = next((p for p in directors if p.is_primary), None) or (
        directors[0] if directors else None
    )

    contact = next((p for p in persons if p.is_primary), None)
    if contact is None:
        contact = next((p for p in persons if p.role == "CONTACT"), None)
    if contact is None and persons:
        contact = persons[0]

    return {
        "director_name": getattr(director, "full_name", "") or "",
        "director_phone": getattr(director, "phone", "") or "",
        "director_email": getattr(director, "email", "") or "",
        "contact_name": getattr(contact, "full_name", "") or "",
        "contact_phone": getattr(contact, "phone", "") or "",
        "contact_email": getattr(contact, "email", "") or "",
    }


def binding_subject_snapshot(binding: ProcessBinding) -> dict:
    if binding.employee_id:
        e = binding.employee
        client = getattr(e, "client_company", None)
        snapshot: dict = {
            "kind": "EMPLOYEE",
            "id": e.id,
            "name": f"{e.first_name} {e.last_name}".strip(),
            "email": e.email or "",
            "employee": {
                "id": e.id,
                "first_name": e.first_name,
                "last_name": e.last_name,
                "full_name": f"{e.first_name} {e.last_name}".strip(),
                "email": e.email or "",
                "org_unit": e.org_unit or "",
                "position": e.position or "",
                "father_name": getattr(e, "father_name", "") or "",
                "national_id": getattr(e, "national_id", "") or "",
                "date_of_birth": (
                    e.date_of_birth.isoformat() if getattr(e, "date_of_birth", None) else ""
                ),
                "place_of_birth": getattr(e, "place_of_birth", "") or "",
                "occupation": getattr(e, "occupation", "") or "",
                "high_risk_position_name": getattr(
                    e, "high_risk_position_name", ""
                )
                or "",
            },
        }
        if client is not None:
            snapshot["client"] = {
                "id": client.id,
                "name": client.name,
                "tax_id": client.tax_id,
                "registration_number": client.registration_number or "",
                "activity_code": getattr(client, "activity_code", "") or "",
                "address": client.address or "",
                "phone": client.phone or "",
                "email": client.email or "",
                "website": client.website or "",
                "risk_assessment_act_name": (
                    _risk_assessment_act_name(client)
                ),
                "risk_assessment_act_date": (
                    client.risk_assessment_act_date.isoformat()
                    if getattr(client, "risk_assessment_act_date", None)
                    else ""
                ),
                **_resolve_client_contacts(client),
            }
        return snapshot

    if binding.equipment_item_id:
        eq = binding.equipment_item
        client = getattr(eq, "client_company", None)
        snapshot = {
            "kind": "EQUIPMENT",
            "id": eq.id,
            "name": eq.name,
            "inventory_number": eq.inventory_number or "",
            "equipment": {
                "id": eq.id,
                "name": eq.name,
                "category": eq.category or "",
                "inventory_number": eq.inventory_number or "",
                "location": eq.location or "",
            },
        }
        if client is not None:
            snapshot["client"] = {
                "id": client.id,
                "name": client.name,
                "tax_id": client.tax_id,
                "registration_number": client.registration_number or "",
                "activity_code": getattr(client, "activity_code", "") or "",
                "address": client.address or "",
                "phone": client.phone or "",
                "email": client.email or "",
                "website": client.website or "",
                "risk_assessment_act_name": (
                    _risk_assessment_act_name(client)
                ),
                "risk_assessment_act_date": (
                    client.risk_assessment_act_date.isoformat()
                    if getattr(client, "risk_assessment_act_date", None)
                    else ""
                ),
                **_resolve_client_contacts(client),
            }
        return snapshot

    if binding.client_company_id:
        c = binding.client_company
        return {
            "kind": "CLIENT_COMPANY",
            "id": c.id,
            "name": c.name,
            "email": c.email or "",
            "client": {
                "id": c.id,
                "name": c.name,
                "tax_id": c.tax_id,
                "registration_number": c.registration_number or "",
                "activity_code": getattr(c, "activity_code", "") or "",
                "address": c.address or "",
                "phone": c.phone or "",
                "email": c.email or "",
                "website": c.website or "",
                "risk_assessment_act_name": (
                    _risk_assessment_act_name(c)
                ),
                "risk_assessment_act_date": (
                    c.risk_assessment_act_date.isoformat()
                    if getattr(c, "risk_assessment_act_date", None)
                    else ""
                ),
                **_resolve_client_contacts(c),
            },
        }

    return {"kind": binding.subject_kind}


def _normalize_snapshot_for_context(snapshot: dict) -> dict:
    if not isinstance(snapshot, dict):
        return snapshot
    out = dict(snapshot)
    emp = out.get("employee")
    if isinstance(emp, dict):
        out["employee"] = normalize_employee_snapshot_national_id(emp)
    for key in ("client",):
        block = out.get(key)
        if isinstance(block, dict):
            out[key] = normalize_client_snapshot_tax_id(block)
    return out


def _build_document_context(run: ProcessRun, snapshot: dict) -> dict:
    ctx = _normalize_snapshot_for_context(snapshot)
    ctx["scheduled_for"] = format_date_display(run.scheduled_for)
    ctx["performed_at"] = format_date_display(run.performed_at)
    ctx["valid_until"] = format_date_display(run.valid_until)
    ctx["process_type_name"] = run.process_type.name if run.process_type_id else ""
    ctx["run_id"] = run.id
    try:
        ctx["instruction_number"] = get_next_code("instruction", "UP")
    except Exception:
        ctx["instruction_number"] = ""
    last_exam_date = ""
    try:
        binding = run.process_binding
        if getattr(binding, "employee_id", None):
            prev_run = (
                ProcessRun.objects.filter(
                    process_binding__employee_id=binding.employee_id,
                    process_type_id=run.process_type_id,
                    status=ProcessRun.STATUS_COMPLETED,
                )
                .exclude(id=run.id)
                .order_by("-performed_at")
                .values_list("performed_at", flat=True)
                .first()
            )
            if prev_run:
                last_exam_date = format_date_display(prev_run)
    except Exception:
        pass
    ctx["last_exam_date"] = last_exam_date
    emp = snapshot.get("employee") if isinstance(
        snapshot.get("employee"), dict) else {}
    dob = emp.get("date_of_birth") or snapshot.get("date_of_birth")
    if dob and isinstance(dob, str) and len(dob) >= 4:
        ctx["year_of_birth"] = dob[:4]
        ctx["date_of_birth"] = format_date_display(dob[:10])
    else:
        ctx["year_of_birth"] = str(dob)[:4] if dob else ""
        ctx["date_of_birth"] = format_date_display(dob) if dob else ""
    snapshot_values = [
        v for k, v in snapshot.items() if k != "kind" and v is not None
    ]
    for i, val in enumerate(snapshot_values, start=1):
        ctx[f"field_{i}"] = str(val)
    return ctx


_UPUT_TEMPLATE_KINDS = {
    "Uput za prethodni lekarski pregled": "prethodni",
    "Uput za lekarski pregled": "periodicni",
}


def _render_template_body(body: str, context: dict) -> str:
    if not body.strip():
        return ""
    t = Template(body)
    return t.render(**{k: (v if v is not None else "") for k, v in context.items()})


def _fill_cell_map_cells(doc: docx.Document, config: dict, context: dict) -> None:
    for entry in config.get("cells") or []:
        table_index = int(entry.get("table", 0))
        row_index = int(entry.get("row", 0))
        col_index = int(entry.get("col", 0))
        field_key = entry.get("fieldKey", "")
        try:
            table = doc.tables[table_index]
            cell = table.rows[row_index].cells[col_index]
        except IndexError:
            continue
        value = _resolve_field_value(field_key, context)
        cell.text = value


def _fill_cell_map_paragraphs(doc: docx.Document, config: dict, context: dict) -> None:
    for entry in config.get("paragraphs") or []:
        contains = entry.get("contains", "")
        field_key = entry.get("fieldKey", "")
        if not contains:
            continue
        try:
            target = next(
                para for para in doc.paragraphs if contains in (para.text or "")
            )
        except StopIteration:
            continue
        value = _resolve_field_value(field_key, context)
        text = target.text or ""
        sep_index = text.find(": ")
        if sep_index != -1:
            new_text = f"{text[:sep_index + 2]}{value}"
        else:
            new_text = f"{text} {value}".strip()
        target.clear()
        target.add_run(new_text)


def apply_cell_map_fill(doc: docx.Document, config: dict, context: dict) -> None:
    _fill_cell_map_cells(doc, config, context)
    _fill_cell_map_paragraphs(doc, config, context)


_PLACEHOLDER_TAG_RE = re.compile(r"\{\{\s*([^{}]+?)\s*\}\}")


def _replace_placeholder_tags(text: str, context: dict) -> str:
    def _sub(match: "re.Match") -> str:
        return _resolve_field_value(match.group(1), context)

    return _PLACEHOLDER_TAG_RE.sub(_sub, text)


def _fill_placeholder_paragraph(para, context: dict) -> None:
    text = para.text or ""
    if "{{" not in text:
        return
    rendered = _replace_placeholder_tags(text, context)
    if not para.runs:
        para.add_run(rendered)
        return
    para.runs[0].text = rendered
    for run in para.runs[1:]:
        run.text = ""


def _iter_table_paragraphs(table):
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                yield para
            for nested_table in cell.tables:
                yield from _iter_table_paragraphs(nested_table)


def _iter_all_paragraphs(doc: docx.Document):
    for para in doc.paragraphs:
        yield para
    for table in doc.tables:
        yield from _iter_table_paragraphs(table)


def apply_placeholder_fill(doc: docx.Document, context: dict) -> None:
    for para in _iter_all_paragraphs(doc):
        _fill_placeholder_paragraph(para, context)


def _row_has_placeholder(row) -> bool:
    return any("{{" in (cell.text or "") for cell in row.cells)


def _find_series_template_row(table):
    for row in table.rows:
        if _row_has_placeholder(row):
            return row
    return None


def _iter_row_paragraphs(row):
    for cell in row.cells:
        for para in cell.paragraphs:
            yield para
        for nested_table in cell.tables:
            yield from _iter_table_paragraphs(nested_table)


def _fill_row_placeholders(row, row_context: dict) -> None:
    for para in _iter_row_paragraphs(row):
        _fill_placeholder_paragraph(para, row_context)


def apply_series_fill(doc: docx.Document, config: dict, context: dict, entity) -> None:
    from .series_sources import get_series_source

    series_list = config.get("series") or []
    for series in series_list:
        try:
            table_index = int(series.get("table", 0))
        except (TypeError, ValueError):
            continue
        try:
            table = doc.tables[table_index]
        except IndexError:
            continue
        template_row = _find_series_template_row(table)
        if template_row is None:
            continue
        pristine_tr = copy.deepcopy(template_row._tr)
        source_fn = get_series_source(series.get("source"))
        items = source_fn(entity, context) if source_fn else []
        anchor_tr = template_row._tr
        for i, item in enumerate(items, start=1):
            new_tr = copy.deepcopy(pristine_tr)
            anchor_tr.addnext(new_tr)
            new_row = _Row(new_tr, table)
            row_context = dict(context)
            row_data = dict(item) if isinstance(item, dict) else {}
            row_data["rbr"] = i
            row_context["r"] = row_data
            _fill_row_placeholders(new_row, row_context)
            anchor_tr = new_tr
        template_row._tr.getparent().remove(template_row._tr)


def apply_docx_fill(
    doc: docx.Document,
    mode: str,
    config: dict,
    context: dict,
    entity=None,
) -> None:
    if mode == "DOCX_PLACEHOLDER":
        if entity is not None and config.get("series"):
            apply_series_fill(doc, config, context, entity)
        apply_placeholder_fill(doc, context)
    else:
        apply_cell_map_fill(doc, config, context)


def _generate_docx_fill(
    run: ProcessRun,
    fill_file,
    snapshot: dict,
    mode: str,
    config: dict,
) -> bytes:
    context = _build_document_context(run, snapshot)
    with fill_file.open("rb") as fh:
        doc = docx.Document(io.BytesIO(fh.read()))
    apply_docx_fill(doc, mode, config, context, entity=run)
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


def _build_company_document_context(company) -> dict:
    return {
        "client": {
            "id": company.id,
            "name": company.name,
            "tax_id": company.tax_id or "",
            "registration_number": company.registration_number or "",
            "address": company.address or "",
            "phone": company.phone or "",
            "email": company.email or "",
        },
    }


def generate_company_document(doc_template, company) -> bytes | None:
    if not doc_template.template_file:
        return None
    name = getattr(doc_template.template_file, "name", "") or ""
    if not name.lower().endswith(".docx"):
        return None
    generation_config = getattr(doc_template, "generation_config", None) or {}
    mode = generation_config.get("mode") or "DOCX_PLACEHOLDER"
    context = _build_company_document_context(company)
    with doc_template.template_file.open("rb") as fh:
        doc = docx.Document(io.BytesIO(fh.read()))
    apply_docx_fill(doc, mode, generation_config, context, entity=company)
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


def _get_system_user():
    return User.objects.filter(is_superuser=True).first()


def _generate_document_for_run(
    run: ProcessRun,
    template: ProcessTemplate,
    snapshot: dict,
) -> DocumentFile | None:
    doc_template = template.document_template
    if not doc_template:
        return None

    existing_prd = (
        ProcessRunDocument.objects.filter(
            process_run=run,
            generated_by_template=template,
        )
        .select_related("document_file")
        .first()
    )
    if existing_prd and existing_prd.document_file_id:
        return existing_prd.document_file

    fill_file = doc_template.template_file
    role_blank_file = None
    blank_placements: list = []
    binding = run.process_binding
    emp = getattr(binding, "employee", None)
    role = getattr(emp, "job_role", None) if emp else None
    if role and getattr(role, "obrazac6_template", None):
        role_blank_file = role.obrazac6_template
        blank_placements = role.obrazac6_fields or []
    elif role and getattr(role, "lzo_revers_template", None):
        role_blank_file = role.lzo_revers_template
        blank_placements = role.lzo_revers_fields or []
    if role_blank_file:
        fill_file = role_blank_file

    system_user = _get_system_user()
    if not system_user:
        logger.warning(
            "Cannot generate document for run id=%s: no system user (superuser) found",
            run.id,
        )
        return None

    category = doc_template.category
    if not category:
        category = DocumentCategory.objects.first()
    if not category:
        logger.warning(
            "Cannot generate document for run id=%s: no document category (template has no category and none in DB)",
            run.id,
        )
        return None

    context = _build_document_context(run, snapshot)
    content_bytes: bytes | None = None
    ext = ".docx"
    title_suffix = f"Run #{run.id}"

    generation_config = getattr(doc_template, "generation_config", None) or {}
    mode = generation_config.get("mode")

    _tname = (doc_template.name or "").strip()
    uput_kind = _UPUT_TEMPLATE_KINDS.get(_tname)
    if uput_kind and emp is not None:
        try:
            from partners.uput import generate_uput
            content_bytes = generate_uput(
                emp, kind=uput_kind, performed_on=run.scheduled_for)
            ext = ".docx"
        except Exception as e:
            logger.warning(
                "Failed to generate uput for run id=%s: %s", run.id, e)

    if not content_bytes and emp is not None:
        try:
            if _tname.startswith("Obrazac 6"):
                from partners.obrazac6 import generate_obrazac6
                content_bytes = generate_obrazac6(
                    emp, context, doc_template.template_file)
                ext = ".docx"
            elif _tname.startswith("Karton zaduženja LZO"):
                from partners.lzo_revers import generate_lzo_revers
                content_bytes = generate_lzo_revers(
                    emp,
                    context=context,
                    template_file=doc_template.template_file,
                )
                ext = ".docx"
        except Exception as e:
            logger.warning(
                "Failed data-driven generation for run id=%s: %s", run.id, e)

    if not content_bytes and mode == "TEMPLATE_BODY" and doc_template.template_body:
        try:
            rendered = _render_template_body(
                doc_template.template_body, context)
            content_bytes = rendered.encode("utf-8")
            ext = ".txt"
        except Exception as e:
            logger.warning(
                "Failed to render template_body for run id=%s: %s",
                run.id,
                e,
            )

    active_file = fill_file
    if not content_bytes and active_file:
        name = getattr(active_file, "name", "") or ""

        if mode == "VISUAL":
            try:
                content_bytes = generate_visual_pdf(
                    doc_template,
                    context,
                    blank_file=role_blank_file,
                    blank_placements=blank_placements,
                    entity=run,
                )
                ext = ".pdf"
            except Exception as e:
                logger.warning(
                    "Failed to generate visual PDF for run id=%s: %s",
                    run.id,
                    e,
                )
        elif mode in ("DOCX_CELL_MAP", "DOCX_PLACEHOLDER"):
            if name and name.lower().endswith(".docx"):
                try:
                    content_bytes = _generate_docx_fill(
                        run,
                        active_file,
                        snapshot,
                        mode,
                        generation_config,
                    )
                    ext = ".docx"
                except Exception as e:
                    logger.warning(
                        "Failed to generate %s docx for run id=%s: %s",
                        mode,
                        run.id,
                        e,
                    )
            else:
                logger.warning(
                    "%s requires a .docx template file for run id=%s template id=%s",
                    mode,
                    run.id,
                    template.id,
                )
        else:
            logger.warning(
                "Unsupported or missing generation mode '%s' for run id=%s template id=%s",
                mode,
                run.id,
                template.id,
            )

    if not content_bytes:
        logger.warning(
            "No content for document generation run id=%s template id=%s (missing template_body and valid template_file)",
            run.id,
            template.id,
        )
        return None

    title = f"{doc_template.name} – {title_suffix}"
    doc_file = DocumentFile(
        category=category,
        title=title,
        uploaded_by=system_user,
        valid_from=run.scheduled_for,
        valid_until=run.valid_until,
    )
    doc_file.file.save(
        f"process_run_{run.id}_{doc_template.id}{ext}",
        ContentFile(content_bytes),
        save=True,
    )
    doc_file.save()

    ProcessRunDocument.objects.create(
        process_run=run,
        document_file=doc_file,
        usage_kind=(
            ProcessRunDocument.USAGE_INVITATION
            if template.trigger
            in (
                ProcessTemplate.TRIGGER_ON_LEAD,
                ProcessTemplate.TRIGGER_ON_SCHEDULED,
            )
            else ProcessRunDocument.USAGE_CERTIFICATE
        ),
        generated_by_template=template,
    )
    logger.info(
        "Generated document id=%s for run id=%s from template id=%s",
        doc_file.id,
        run.id,
        template.id,
    )
    return doc_file


def _resolve_email_recipients(
    template: ProcessTemplate,
    binding: ProcessBinding,
) -> list[str]:
    if template.email_to_kind == ProcessTemplate.EMAIL_TO_CUSTOM:
        addr = (template.custom_email_recipient or "").strip()
        return [addr] if addr else []
    if template.email_to_kind == ProcessTemplate.EMAIL_TO_CLIENT_AND_MAK:
        return resolve_reminder_recipients(binding)
    if template.email_to_kind == ProcessTemplate.EMAIL_TO_MAK:
        return internal_mak_recipients()
    if template.email_to_kind == ProcessTemplate.EMAIL_TO_CLIENT_MAIN:
        company = _binding_company(binding)
        addr = (company.email or "").strip() if company else ""
        return [addr] if addr else []
    if (
        template.email_to_kind == ProcessTemplate.EMAIL_TO_EMPLOYEE
        and binding.employee_id
    ):
        addr = (binding.employee.email or "").strip()
        return [addr] if addr else []
    if template.email_to_kind == ProcessTemplate.EMAIL_TO_INTERNAL_ROLE:
        group_id = template.notification_role_group_id
        if not group_id:
            logger.warning(
                "INTERNAL_ROLE email: no notification_role_group on ProcessTemplate id=%s",
                template.id,
            )
            return []
        emails = list(
            User.objects.filter(
                groups__id=group_id,
                is_active=True,
            )
            .exclude(email="")
            .values_list("email", flat=True)
            .distinct()
        )
        cleaned = [e.strip() for e in emails if e and e.strip()]
        if not cleaned:
            logger.warning(
                "INTERNAL_ROLE email: no active users with email in group id=%s (template id=%s)",
                group_id,
                template.id,
            )
        return cleaned
    return []


def _binding_company(binding: ProcessBinding):
    if binding.client_company_id:
        return binding.client_company
    if binding.employee_id and binding.employee.client_company_id:
        return binding.employee.client_company
    if binding.equipment_item_id and binding.equipment_item.client_company_id:
        return binding.equipment_item.client_company
    return None


def internal_mak_recipients() -> list[str]:
    group_name = getattr(settings, "REMINDER_INTERNAL_GROUP", "") or ""
    internal = None
    if group_name:
        internal = User.objects.filter(
            groups__name=group_name,
            is_active=True,
        ).exclude(email="")
    if internal is None or not internal.exists():
        internal = User.objects.filter(is_active=True, is_staff=True).exclude(
            email=""
        )
    recipients: list[str] = []
    for addr in internal.values_list("email", flat=True).distinct():
        addr = (addr or "").strip()
        if addr and addr not in recipients:
            recipients.append(addr)
    return recipients


def resolve_reminder_recipients(binding: ProcessBinding) -> list[str]:
    recipients: list[str] = []
    company = _binding_company(binding)
    if company is not None:
        addr = (company.email or "").strip()
        if addr:
            recipients.append(addr)
    for addr in internal_mak_recipients():
        if addr not in recipients:
            recipients.append(addr)
    return recipients


def build_generic_reminder(run: ProcessRun, offset_days: int) -> tuple[str, str]:
    snapshot = run.subject_snapshot or {}
    subject_name = snapshot.get("name") or snapshot.get("kind") or ""
    name = run.process_type.name if run.process_type_id else "Obaveza"
    due = format_date_display(run.scheduled_for)
    if offset_days < 0:
        subject = f"Podsetnik: {name}"
        line = f"Obaveza „{name}” za {subject_name} dospeva {due}."
    elif offset_days == 0:
        subject = f"Danas je rok: {name}"
        line = f"Danas ({due}) je rok za obavezu „{name}” za {subject_name}."
    else:
        subject = f"Prekoračen rok: {name}"
        line = f"Rok za obavezu „{name}” za {subject_name} je istekao {due}."
    body = f"{line}\n\nMolimo da se obaveza izvrši i evidentira u aplikaciji."
    return subject.strip(), body


def send_generic_reminder(
    run: ProcessRun,
    binding: ProcessBinding,
    offset_days: int,
    *,
    fail_silently: bool = True,
) -> tuple[bool, list[str], str, str]:
    recipients = resolve_reminder_recipients(binding)
    subject, body = build_generic_reminder(run, offset_days)
    if not recipients:
        if not fail_silently:
            raise ValueError(
                "Nema primaoca za podsetnik (firma bez mejla i nema internih korisnika)."
            )
        return False, [], subject, body
    with company_email_test_mode(_binding_company(binding)):
        sent = get_email_sender().send(
            recipients=recipients,
            subject=subject,
            body=body,
            attachments=None,
            fail_silently=fail_silently,
        )
    if not sent and not fail_silently:
        raise RuntimeError("Slanje podsetnika nije uspelo.")
    return bool(sent), recipients, subject, body


def _read_document_file_attachment(
    doc_file: DocumentFile,
) -> tuple[str, bytes] | None:
    if not doc_file.file:
        return None
    try:
        with doc_file.file.open("rb") as fh:
            data = fh.read()
    except Exception as exc:
        logger.warning(
            "Could not read attachment for doc id=%s: %s",
            doc_file.id,
            exc,
        )
        return None
    name = doc_file.file.name.split("/")[-1]
    return name, data


def _collect_template_email_attachments(
    run: ProcessRun,
    template: ProcessTemplate,
    generated_document: DocumentFile | None = None,
) -> list[tuple[str, bytes]]:
    attachments: list[tuple[str, bytes]] = []
    seen_ids: set[int] = set()

    def add_doc_file(doc_file: DocumentFile | None) -> None:
        if doc_file is None or doc_file.id in seen_ids:
            return
        item = _read_document_file_attachment(doc_file)
        if item is None:
            return
        seen_ids.add(doc_file.id)
        attachments.append(item)

    if template.attach_generated_document and generated_document is not None:
        add_doc_file(generated_document)

    if template.attach_uploaded_documents:
        uploaded = run.documents.select_related("document_file").filter(
            generated_by_template__isnull=True,
        )
        for prd in uploaded:
            add_doc_file(prd.document_file)

    return attachments


def _send_email_for_template(
    template: ProcessTemplate,
    binding: ProcessBinding,
    snapshot: dict,
    run: ProcessRun | None = None,
    generated_document: DocumentFile | None = None,
    *,
    fail_silently: bool = True,
) -> bool:
    recipients = _resolve_email_recipients(template, binding)
    if not recipients:
        logger.warning(
            "No email recipient for ProcessTemplate id=%s",
            template.id,
        )
        if not fail_silently:
            raise ValueError("Nema podešenog primaoca mejla.")
        return False

    subject = template.email_subject_template or "Process notification"
    body = template.email_body_template or ""
    if run is not None:
        context = _build_document_context(run, snapshot)
        subject = _render_template_body(subject, context)
        body = _render_template_body(body, context)
    attachments = None
    if run is not None:
        collected = _collect_template_email_attachments(
            run,
            template,
            generated_document,
        )
        attachments = collected if collected else None
    try:
        with company_email_test_mode(_binding_company(binding)):
            sent = get_email_sender().send(
                recipients=recipients,
                subject=subject,
                body=body,
                attachments=attachments,
                fail_silently=fail_silently,
            )
        if not sent:
            logger.warning(
                "Email send returned False for ProcessTemplate id=%s",
                template.id,
            )
            if not fail_silently:
                raise RuntimeError("Email sender returned False")
            return False
        return True
    except Exception as e:
        logger.exception(
            "Failed to send email for ProcessTemplate id=%s: %s",
            template.id,
            e,
        )
        if not fail_silently:
            raise
        return False


def execute_template_actions(
    trigger: str,
    run: ProcessRun,
    binding: ProcessBinding,
    snapshot: dict,
    template: ProcessTemplate,
) -> tuple[DocumentFile | None, bool, str]:
    logger.info(
        "%s trigger: run_id=%s process_type=%s template_id=%s template=%s",
        trigger,
        run.id,
        run.process_type_id,
        template.id,
        template,
    )

    generated_document: DocumentFile | None = None
    if template.generate_document and template.document_template_id:
        try:
            generated_document = _generate_document_for_run(
                run, template, snapshot,
            )
        except Exception as e:
            logger.exception(
                "Failed to generate document in %s for run id=%s template id=%s: %s",
                trigger,
                run.id,
                template.id,
                e,
            )

    email_sent = False
    email_error = ""
    if template.send_email:
        try:
            email_sent = _send_email_for_template(
                template,
                binding,
                snapshot,
                run=run,
                generated_document=generated_document,
                fail_silently=False,
            )
        except Exception as e:
            email_error = str(e)
            logger.exception(
                "Failed to send email in %s for run id=%s template id=%s: %s",
                trigger,
                run.id,
                template.id,
                e,
            )
        if not email_sent and not email_error:
            email_error = "Mejl nije poslat."
    return generated_document, email_sent, email_error
