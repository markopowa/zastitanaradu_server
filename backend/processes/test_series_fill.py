import io
import uuid
from datetime import date
from unittest.mock import patch

import docx
from django.contrib.auth import get_user_model
from django.core.files.base import ContentFile
from django.test import TestCase
from rest_framework.test import APIClient

from documents.models import DocumentCategory, DocumentTemplate
from partners.models import ClientCompany, Employee
from processes.models import ProcessBinding, ProcessRun, ProcessType
from processes.utils import apply_series_fill, generate_company_document

User = get_user_model()


def make_company(**kwargs):
    defaults = {"name": "Test firma", "tax_id": uuid.uuid4().hex[:9]}
    defaults.update(kwargs)
    return ClientCompany.objects.create(**defaults)


def make_process_type(**kwargs):
    defaults = {
        "code": f"SERIES_TEST_{uuid.uuid4().hex[:12]}",
        "name": "Test obaveza",
        "subject_kind": ProcessType.SUBJECT_EMPLOYEE,
        "domain": ProcessType.DOMAIN_BZNR,
        "shape": ProcessType.SHAPE_PERIODIC,
        "proof_kind": ProcessType.PROOF_RECORDED,
        "applicability_rule": {"always": True},
        "include_in_medical_exam_record": True,
        "is_active": True,
    }
    defaults.update(kwargs)
    return ProcessType.objects.create(**defaults)


def build_series_docx(extra_paragraph: str = "") -> bytes:
    doc = docx.Document()
    if extra_paragraph:
        doc.add_paragraph(extra_paragraph)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Redni broj"
    table.cell(0, 1).text = "Ime"
    table.cell(1, 0).text = "{{r.rbr}}"
    table.cell(1, 1).text = "{{r.name}}"
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


def load_doc(content_bytes: bytes) -> docx.Document:
    return docx.Document(io.BytesIO(content_bytes))


def table_rows_text(doc: docx.Document, table_index: int = 0) -> list[list[str]]:
    table = doc.tables[table_index]
    return [[cell.text for cell in row.cells] for row in table.rows]


def row_has_placeholder(row_texts: list[str]) -> bool:
    return any("{{" in text for text in row_texts)


class SeriesFillEngineTest(TestCase):
    def _config(self, source_key="test_source", table=0):
        return {"series": [{"source": source_key, "table": table}]}

    def test_row_multiplication_and_template_row_removed(self):
        doc = docx.Document(io.BytesIO(build_series_docx()))
        items = [{"name": "Ana"}, {"name": "Bora"}, {"name": "Vera"}]
        with patch.dict(
            "processes.series_sources.SERIES_SOURCES",
            {"test_source": lambda entity, context: items},
        ):
            apply_series_fill(doc, self._config(), {}, entity=object())

        table = doc.tables[0]
        self.assertEqual(len(table.rows), 4)
        rows = table_rows_text(doc)
        self.assertFalse(row_has_placeholder(rows[0]))
        for row in rows[1:]:
            self.assertFalse(row_has_placeholder(row))
        self.assertEqual([r[1] for r in rows[1:]], ["Ana", "Bora", "Vera"])

    def test_rbr_numbering(self):
        doc = docx.Document(io.BytesIO(build_series_docx()))
        items = [{"name": "Ana"}, {"name": "Bora"}, {"name": "Vera"}]
        with patch.dict(
            "processes.series_sources.SERIES_SOURCES",
            {"test_source": lambda entity, context: items},
        ):
            apply_series_fill(doc, self._config(), {}, entity=object())

        rows = table_rows_text(doc)
        self.assertEqual([r[0] for r in rows[1:]], ["1", "2", "3"])

    def test_empty_source_removes_template_row_with_no_replacement(self):
        doc = docx.Document(io.BytesIO(build_series_docx()))
        with patch.dict(
            "processes.series_sources.SERIES_SOURCES",
            {"test_source": lambda entity, context: []},
        ):
            apply_series_fill(doc, self._config(), {}, entity=object())

        table = doc.tables[0]
        self.assertEqual(len(table.rows), 1)
        rows = table_rows_text(doc)
        self.assertFalse(row_has_placeholder(rows[0]))

    def test_combined_series_and_scalar_tags(self):
        content = build_series_docx(extra_paragraph="Firma: {{client.name}}")
        doc = docx.Document(io.BytesIO(content))
        items = [{"name": "Ana"}, {"name": "Bora"}]
        context = {"client": {"name": "Acme d.o.o."}}
        with patch.dict(
            "processes.series_sources.SERIES_SOURCES",
            {"test_source": lambda entity, context: items},
        ):
            from processes.utils import apply_placeholder_fill

            apply_series_fill(doc, self._config(), context, entity=object())
            apply_placeholder_fill(doc, context)

        self.assertEqual(doc.paragraphs[0].text, "Firma: Acme d.o.o.")
        rows = table_rows_text(doc)
        self.assertEqual([r[1] for r in rows[1:]], ["Ana", "Bora"])

    def test_missing_table_index_is_noop(self):
        doc = docx.Document(io.BytesIO(build_series_docx()))
        config = {"series": [{"source": "test_source", "table": 5}]}
        with patch.dict(
            "processes.series_sources.SERIES_SOURCES",
            {"test_source": lambda entity, context: [{"name": "Ana"}]},
        ):
            apply_series_fill(doc, config, {}, entity=object())
        self.assertEqual(len(doc.tables[0].rows), 2)

    def test_unknown_source_key_removes_template_row(self):
        doc = docx.Document(io.BytesIO(build_series_docx()))
        config = {"series": [{"source": "does_not_exist", "table": 0}]}
        apply_series_fill(doc, config, {}, entity=object())
        self.assertEqual(len(doc.tables[0].rows), 1)


class Obrazac1ViaTemplateEndpointTest(TestCase):
    def setUp(self):
        self.user = User.objects.create_superuser(
            username=f"admin-{uuid.uuid4().hex[:8]}",
            password="x",
            email=f"admin-{uuid.uuid4().hex[:8]}@test.local",
        )
        self.client_api = APIClient()
        self.client_api.force_authenticate(user=self.user)
        self.company = make_company(name="Firma Serija")
        self.category = DocumentCategory.objects.create(
            name=f"Kategorija {uuid.uuid4().hex[:8]}")
        self.process_type = make_process_type()

    def _make_obrazac1_template(self):
        content = build_obrazac1_test_docx()
        tpl = DocumentTemplate.objects.create(
            name="Obrazac 1 — evidencija lekarskih pregleda",
            context_type=DocumentTemplate.CONTEXT_CLIENT_COMPANY,
            category=self.category,
            generation_config={
                "mode": "DOCX_PLACEHOLDER",
                "series": [
                    {"source": "completed_medical_exams_for_company", "table": 0},
                ],
            },
        )
        tpl.template_file.save(
            "obrazac1_test.docx", ContentFile(content), save=True,
        )
        return tpl

    def _make_completed_run(self, first_name, last_name):
        employee = Employee.objects.create(
            client_company=self.company,
            first_name=first_name,
            last_name=last_name,
            high_risk_position_name="Rukovalac masinom",
        )
        binding = ProcessBinding.objects.create(
            process_type=self.process_type,
            subject_kind=ProcessBinding.SUBJECT_EMPLOYEE,
            employee=employee,
            is_active=True,
            next_run_at=date.today(),
        )
        ProcessRun.objects.create(
            process_binding=binding,
            process_type=self.process_type,
            scheduled_for=date.today(),
            performed_at=date.today(),
            status=ProcessRun.STATUS_COMPLETED,
            result_data={"report_number": "12/2026"},
        )
        return employee

    def test_generate_company_document_fills_series_rows(self):
        tpl = self._make_obrazac1_template()
        self._make_completed_run("Jovan", "Jovanovic")
        self._make_completed_run("Milica", "Milic")

        content = generate_company_document(tpl, self.company)
        self.assertIsNotNone(content)
        doc = load_doc(content)
        rows = table_rows_text(doc)
        self.assertFalse(row_has_placeholder(rows[0]))
        names = [row[2] for row in rows[1:]]
        self.assertIn("Jovan Jovanovic", names)
        self.assertIn("Milica Milic", names)

    def test_endpoint_uses_template_and_contains_employee_names(self):
        self._make_obrazac1_template()
        self._make_completed_run("Petar", "Petrovic")

        response = self.client_api.get(
            f"/api/partners/client-companies/{self.company.id}/medical-exam-record/",
        )
        self.assertEqual(response.status_code, 200)
        doc = load_doc(response.content)
        all_text = "\n".join(
            cell.text
            for table in doc.tables
            for row in table.rows
            for cell in row.cells
        )
        self.assertIn("Petar Petrovic", all_text)
        self.assertNotIn("{{", all_text)


def build_obrazac1_test_docx() -> bytes:
    doc = docx.Document()
    doc.add_paragraph("{{client.name}}")
    table = doc.add_table(rows=2, cols=3)
    table.cell(0, 0).text = "Redni broj"
    table.cell(0, 1).text = "Naziv radnog mesta"
    table.cell(0, 2).text = "Ime i prezime"
    table.cell(1, 0).text = "{{r.rbr}}"
    table.cell(1, 1).text = "{{r.job_role_name}}"
    table.cell(1, 2).text = "{{r.employee_name}}"
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()
