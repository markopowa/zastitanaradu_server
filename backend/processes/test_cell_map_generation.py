import io
import uuid
from datetime import date

import docx
from django.contrib.auth import get_user_model
from django.core.files.base import ContentFile
from django.test import TestCase

from documents.models import DocumentCategory, DocumentTemplate
from partners.models import ClientCompany, Employee, JobRole
from processes.models import ProcessBinding, ProcessRun, ProcessTemplate, ProcessType
from processes.utils import _generate_document_for_run

User = get_user_model()


def make_company(**kwargs):
    defaults = {"name": "Test firma", "tax_id": uuid.uuid4().hex[:9]}
    defaults.update(kwargs)
    return ClientCompany.objects.create(**defaults)


def make_process_type(**kwargs):
    defaults = {
        "name": "Obrazac 6 test",
        "subject_kind": ProcessType.SUBJECT_EMPLOYEE,
        "domain": ProcessType.DOMAIN_BZNR,
        "shape": ProcessType.SHAPE_PERIODIC,
        "proof_kind": ProcessType.PROOF_RECORDED,
        "applicability_rule": {"always": True},
        "is_active": True,
    }
    defaults.update(kwargs)
    return ProcessType.objects.create(**defaults)


def build_blank_docx_bytes(marker_row1_col2: str = "") -> bytes:
    doc = docx.Document()
    table = doc.add_table(rows=17, cols=4)
    for row_index in range(17):
        for col_index in range(4):
            cell = table.cell(row_index, col_index)
            if row_index == 0 and col_index == 2:
                cell.text = ""
            elif row_index == 6:
                cell.text = ""
            elif row_index == 1 and col_index == 2 and marker_row1_col2:
                cell.text = marker_row1_col2
            else:
                cell.text = f"label_{row_index}_{col_index}"
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


def read_table_text(content_bytes: bytes, table_index: int, row: int, col: int) -> str:
    doc = docx.Document(io.BytesIO(content_bytes))
    return doc.tables[table_index].rows[row].cells[col].text


def docx_contains(content_bytes: bytes, needle: str) -> bool:
    doc = docx.Document(io.BytesIO(content_bytes))
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if needle in cell.text:
                    return True
    return False


class CellMapGenerationTest(TestCase):
    def setUp(self):
        self.user = User.objects.create_superuser(
            username="cellmapgen", password="x", email="cellmapgen@test.local"
        )
        self.category = DocumentCategory.objects.create(name="Obrazac 6")
        self.doc_tpl = DocumentTemplate.objects.create(
            name="Obrazac 6",
            context_type=DocumentTemplate.CONTEXT_EMPLOYEE,
            generation_config={
                "mode": "DOCX_CELL_MAP",
                "cells": [
                    {"table": 0, "row": 0, "col": 2, "fieldKey": "employee.full_name"},
                    {"table": 0, "row": 6, "col": 0, "fieldKey": "performed_at"},
                ],
            },
            category=self.category,
        )
        self.doc_tpl.template_file.save(
            "obrazac6_default.docx",
            ContentFile(build_blank_docx_bytes()),
            save=True,
        )
        self.company = make_company(name="Firma Obrazac6")
        self.process_type = make_process_type()

    def _make_binding_and_run(self, employee):
        binding = ProcessBinding.objects.create(
            process_type=self.process_type,
            subject_kind=ProcessBinding.SUBJECT_EMPLOYEE,
            employee=employee,
            is_active=True,
            next_run_at=date.today(),
        )
        run = ProcessRun.objects.create(
            process_binding=binding,
            process_type=self.process_type,
            scheduled_for=date.today(),
            performed_at=date.today(),
            status=ProcessRun.STATUS_PENDING,
            subject_snapshot={
                "kind": "EMPLOYEE",
                "employee": {
                    "first_name": employee.first_name,
                    "last_name": employee.last_name,
                    "full_name": f"{employee.first_name} {employee.last_name}".strip(),
                    "email": "",
                    "org_unit": "",
                    "position": "",
                    "father_name": "",
                    "national_id": "",
                    "date_of_birth": "",
                    "place_of_birth": "",
                    "occupation": "",
                    "high_risk_position_name": "",
                },
                "client": {
                    "name": self.company.name,
                    "tax_id": self.company.tax_id,
                    "registration_number": "",
                    "activity_code": "",
                    "address": "",
                    "phone": "",
                    "email": "",
                    "website": "",
                    "risk_assessment_act_name": "",
                    "risk_assessment_act_date": "",
                },
            },
        )
        process_tpl = ProcessTemplate.objects.create(
            process_type=self.process_type,
            document_template=self.doc_tpl,
            trigger=ProcessTemplate.TRIGGER_ON_SCHEDULED,
            generate_document=True,
            send_email=False,
        )
        return binding, run, process_tpl

    def test_role_blank_used_when_employee_has_job_role(self):
        role = JobRole.objects.create(
            client_company=self.company,
            name="Varilac",
        )
        role.obrazac6_template.save(
            "obrazac6_varilac.docx",
            ContentFile(build_blank_docx_bytes(marker_row1_col2="ROLE_MARKER")),
            save=True,
        )
        employee = Employee.objects.create(
            client_company=self.company,
            first_name="Marko",
            last_name="Markovic",
            job_role=role,
        )
        _binding, run, process_tpl = self._make_binding_and_run(employee)
        result = _generate_document_for_run(
            run, process_tpl, run.subject_snapshot,
        )
        self.assertIsNotNone(result)
        with result.file.open("rb") as fh:
            content = fh.read()
        self.assertEqual(read_table_text(content, 0, 0, 2), "Marko Markovic")
        self.assertTrue(docx_contains(content, "ROLE_MARKER"))

    def test_falls_back_to_document_template_file_without_job_role(self):
        employee = Employee.objects.create(
            client_company=self.company,
            first_name="Ana",
            last_name="Anic",
        )
        _binding, run, process_tpl = self._make_binding_and_run(employee)
        result = _generate_document_for_run(
            run, process_tpl, run.subject_snapshot,
        )
        self.assertIsNotNone(result)
        with result.file.open("rb") as fh:
            content = fh.read()
        self.assertEqual(read_table_text(content, 0, 0, 2), "Ana Anic")
        self.assertFalse(docx_contains(content, "ROLE_MARKER"))
