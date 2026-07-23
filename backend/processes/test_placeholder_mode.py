import io
import uuid
from datetime import date

import docx
from django.contrib.auth import get_user_model
from django.core.files.base import ContentFile
from django.test import TestCase

from documents.models import DocumentCategory, DocumentTemplate
from partners.document_generation import generate_employee_document
from partners.models import ClientCompany, Employee, JobRole
from processes.models import ProcessBinding, ProcessRun, ProcessTemplate, ProcessType
from processes.utils import _generate_document_for_run

User = get_user_model()


def make_company(**kwargs):
    defaults = {"name": "Test firma", "tax_id": uuid.uuid4().hex[:9]}
    defaults.update(kwargs)
    return ClientCompany.objects.create(**defaults)


def make_process_type(code=None, **kwargs):
    defaults = {
        "code": code or f"TEST_{uuid.uuid4().hex[:12]}",
        "name": "Proces test",
        "subject_kind": ProcessType.SUBJECT_EMPLOYEE,
        "domain": ProcessType.DOMAIN_BZNR,
        "shape": ProcessType.SHAPE_PERIODIC,
        "proof_kind": ProcessType.PROOF_RECORDED,
        "applicability_rule": {"always": True},
        "is_active": True,
    }
    defaults.update(kwargs)
    return ProcessType.objects.create(**defaults)


def build_paragraph_tag_docx(text: str) -> bytes:
    doc = docx.Document()
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


def build_table_tag_docx(cell_text: str) -> bytes:
    doc = docx.Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "label"
    table.cell(1, 1).text = cell_text
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


def build_split_run_docx(prefix: str, suffix: str) -> bytes:
    doc = docx.Document()
    para = doc.add_paragraph()
    para.add_run(prefix)
    para.add_run(suffix)
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


def paragraph_texts(content_bytes: bytes) -> list[str]:
    doc = docx.Document(io.BytesIO(content_bytes))
    return [p.text for p in doc.paragraphs]


def all_cell_texts(content_bytes: bytes) -> list[str]:
    doc = docx.Document(io.BytesIO(content_bytes))
    texts = []
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                texts.append(cell.text)
    return texts


class PlaceholderModeGenerationTest(TestCase):
    def setUp(self):
        self.user = User.objects.create_superuser(
            username="placeholdergen", password="x", email="placeholdergen@test.local",
        )
        self.category = DocumentCategory.objects.create(
            name="Placeholder test kategorija")
        self.company = make_company(name="Firma Placeholder")
        self.process_type = make_process_type(code="PLACEHOLDER_TEST")

    def _make_doc_template(
        self, content_bytes: bytes, name: str = "Placeholder sablon",
    ) -> DocumentTemplate:
        doc_tpl = DocumentTemplate.objects.create(
            name=name,
            context_type=DocumentTemplate.CONTEXT_EMPLOYEE,
            category=self.category,
            generation_config={"mode": "DOCX_PLACEHOLDER"},
        )
        doc_tpl.template_file.save(
            "placeholder_template.docx",
            ContentFile(content_bytes),
            save=True,
        )
        return doc_tpl

    def _make_binding_and_run(self, employee, process_type=None):
        process_type = process_type or self.process_type
        binding = ProcessBinding.objects.create(
            process_type=process_type,
            subject_kind=ProcessBinding.SUBJECT_EMPLOYEE,
            employee=employee,
            is_active=True,
            next_run_at=date.today(),
        )
        run = ProcessRun.objects.create(
            process_binding=binding,
            process_type=process_type,
            scheduled_for=date.today(),
            performed_at=date.today(),
            status=ProcessRun.STATUS_PENDING,
            subject_snapshot={
                "kind": "EMPLOYEE",
                "employee": {
                    "first_name": employee.first_name,
                    "last_name": employee.last_name,
                    "full_name": f"{employee.first_name} {employee.last_name}".strip(),
                    "email": "",
                    "org_unit": "",
                    "position": "",
                    "father_name": "",
                    "national_id": "",
                    "date_of_birth": "",
                    "place_of_birth": "",
                    "occupation": "",
                    "high_risk_position_name": "",
                },
                "client": {
                    "name": self.company.name,
                    "tax_id": self.company.tax_id,
                    "registration_number": "",
                    "activity_code": "",
                    "address": "",
                    "phone": "",
                    "email": "",
                    "website": "",
                    "risk_assessment_act_name": "",
                    "risk_assessment_act_date": "",
                },
            },
        )
        return binding, run

    def _generate(self, employee, doc_tpl, process_type=None):
        _binding, run = self._make_binding_and_run(employee, process_type)
        process_tpl = ProcessTemplate.objects.create(
            process_type=process_type or self.process_type,
            document_template=doc_tpl,
            trigger=ProcessTemplate.TRIGGER_ON_SCHEDULED,
            generate_document=True,
            send_email=False,
        )
        result = _generate_document_for_run(
            run, process_tpl, run.subject_snapshot,
        )
        self.assertIsNotNone(result)
        with result.file.open("rb") as fh:
            return fh.read()

    def test_tag_in_paragraph_is_replaced(self):
        doc_tpl = self._make_doc_template(
            build_paragraph_tag_docx("Zaposleni: {{ employee.full_name }}"),
        )
        employee = Employee.objects.create(
            client_company=self.company, first_name="Marko", last_name="Markovic",
        )
        content = self._generate(employee, doc_tpl)
        self.assertIn("Zaposleni: Marko Markovic", paragraph_texts(content))

    def test_tag_in_table_cell_is_replaced(self):
        doc_tpl = self._make_doc_template(
            build_table_tag_docx("{{ client.name }}"),
        )
        employee = Employee.objects.create(
            client_company=self.company, first_name="Ana", last_name="Anic",
        )
        content = self._generate(employee, doc_tpl)
        self.assertIn(self.company.name, all_cell_texts(content))

    def test_tag_split_across_runs_is_replaced(self):
        doc_tpl = self._make_doc_template(
            build_split_run_docx("Ime: {{ employee.full_name", " }}"),
        )
        employee = Employee.objects.create(
            client_company=self.company, first_name="Petar", last_name="Petrovic",
        )
        content = self._generate(employee, doc_tpl)
        self.assertIn("Ime: Petar Petrovic", paragraph_texts(content))

    def test_spaced_tag_is_replaced(self):
        doc_tpl = self._make_doc_template(
            build_paragraph_tag_docx("Firma: {{   client.name   }}"),
        )
        employee = Employee.objects.create(
            client_company=self.company, first_name="Jovan", last_name="Jovanovic",
        )
        content = self._generate(employee, doc_tpl)
        self.assertIn(f"Firma: {self.company.name}", paragraph_texts(content))

    def test_unresolved_tag_becomes_empty(self):
        doc_tpl = self._make_doc_template(
            build_paragraph_tag_docx(
                "Nepoznato: {{ employee.nonexistent_field }}."),
        )
        employee = Employee.objects.create(
            client_company=self.company, first_name="Nikola", last_name="Nikolic",
        )
        content = self._generate(employee, doc_tpl)
        self.assertIn("Nepoznato: .", paragraph_texts(content))

    def test_role_blank_override_wins_over_template_file(self):
        doc_tpl = self._make_doc_template(
            build_paragraph_tag_docx(
                "Podrazumevano: {{ employee.full_name }}"),
        )
        role = JobRole.objects.create(
            client_company=self.company, name="Varilac")
        role.obrazac6_template.save(
            "role_blank.docx",
            ContentFile(build_paragraph_tag_docx(
                "ROLE_BLANK: {{ employee.full_name }}")),
            save=True,
        )
        employee = Employee.objects.create(
            client_company=self.company, first_name="Stevan", last_name="Stevanovic",
            job_role=role,
        )
        content = self._generate(employee, doc_tpl)
        texts = paragraph_texts(content)
        self.assertIn("ROLE_BLANK: Stevan Stevanovic", texts)
        self.assertNotIn("Podrazumevano: Stevan Stevanovic", texts)

    def test_on_demand_obrazac6_is_data_driven_pdf(self):
        doc_tpl = DocumentTemplate.objects.create(
            name="Obrazac 6 — evidencija o osposobljenosti za bezbedan rad",
            context_type=DocumentTemplate.CONTEXT_EMPLOYEE,
            category=self.category,
            generation_config={"mode": "VISUAL", "placeholders": []},
        )
        doc_tpl.template_file.save(
            "obrazac6_blank.docx",
            ContentFile(build_paragraph_tag_docx("Obrazac 6")),
            save=True,
        )
        obrazac6_pt = make_process_type(code="OSPOSOBLJAVANJE_BZR")
        role = JobRole.objects.create(
            client_company=self.company, name="Elektricar")
        employee = Employee.objects.create(
            client_company=self.company, first_name="Dragan", last_name="Draganovic",
            job_role=role,
        )
        self._make_binding_and_run(employee, obrazac6_pt)

        _doc_file, content_bytes = generate_employee_document(
            employee, "OBRAZAC6")
        self.assertTrue(content_bytes.startswith(b"%PDF"))

    def test_uput_generation_substitutes_client_and_employee_values(self):
        uput_body = (
            "UPUT ZA PERIODICNI LEKARSKI PREGLED\n"
            "Poslodavac: {{ client.name }}\n"
            "Zaposleni: {{ employee.first_name }} {{ employee.last_name }}\n"
        )
        uput_doc_tpl = self._make_doc_template(
            build_paragraph_tag_docx(uput_body),
            name="Uput za lekarski pregled test",
        )
        lekarski_pt = make_process_type(code="UPUT_TEST")
        employee = Employee.objects.create(
            client_company=self.company, first_name="Milan", last_name="Milic",
        )
        content = self._generate(employee, uput_doc_tpl, lekarski_pt)
        texts = paragraph_texts(content)
        self.assertTrue(
            any(f"Poslodavac: {self.company.name}" in t for t in texts))
        self.assertTrue(any("Zaposleni: Milan Milic" in t for t in texts))
