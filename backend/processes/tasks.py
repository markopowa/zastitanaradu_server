from datetime import date, timedelta
import io
import logging
import docx
from django.conf import settings
from django.contrib.auth import get_user_model
from django.core.files.base import ContentFile
from core.email_sender import get_email_sender
from jinja2 import Template

from documents.models import DocumentCategory, DocumentFile, DocumentTemplate

from .models import ProcessBinding, ProcessRun, ProcessRunDocument, ProcessTemplate
from .utils import binding_subject_snapshot

logger = logging.getLogger(__name__)
User = get_user_model()


def run_process_binding(binding_id: int) -> None:
    try:
        binding = (
            ProcessBinding.objects.select_related(
                "process_type", "employee", "equipment_item", "client_company"
            )
            .get(id=binding_id)
        )
    except ProcessBinding.DoesNotExist:
        logger.warning("ProcessBinding id=%s not found", binding_id)
        return

    if not binding.is_active:
        logger.info("ProcessBinding id=%s is inactive, skipping", binding_id)
        return

    pt = binding.process_type
    scheduled_for = binding.next_run_at or date.today()
    snapshot = binding_subject_snapshot(binding)

    run = ProcessRun.objects.create(
        process_binding=binding,
        process_type=pt,
        subject_snapshot=snapshot,
        scheduled_for=scheduled_for,
        status=ProcessRun.STATUS_PENDING,
    )

    templates = ProcessTemplate.objects.filter(
        process_type=pt,
        trigger=ProcessTemplate.TRIGGER_ON_SCHEDULED,
    ).select_related("document_template")

    for template in templates:
        if template.generate_document and template.document_template_id:
            _generate_document_for_run(run, template, snapshot)
        if template.send_email:
            _send_email_for_template(template, binding, snapshot, run=run)

    period_months = binding.custom_period_months or pt.default_period_months or 12
    performed_at = scheduled_for
    valid_until = performed_at + timedelta(days=period_months * 30)

    run.performed_at = performed_at
    run.valid_until = valid_until
    run.status = ProcessRun.STATUS_COMPLETED
    run.save(update_fields=["performed_at", "valid_until", "status"])

    binding.last_run_at = performed_at
    binding.next_run_at = valid_until
    binding.save(update_fields=["last_run_at", "next_run_at"])

    logger.info(
        "ProcessBinding id=%s run id=%s completed, valid_until=%s next_run_at=%s",
        binding_id,
        run.id,
        valid_until,
        binding.next_run_at,
    )


def _build_document_context(run: ProcessRun, snapshot: dict) -> dict:
    """Build context dict for template rendering from run and subject snapshot."""
    ctx = dict(snapshot)
    ctx["scheduled_for"] = str(run.scheduled_for) if run.scheduled_for else ""
    ctx["performed_at"] = str(run.performed_at) if run.performed_at else ""
    ctx["valid_until"] = str(run.valid_until) if run.valid_until else ""
    ctx["process_type_name"] = run.process_type.name if run.process_type_id else ""
    ctx["run_id"] = run.id
    # Backward compatibility: field_1, field_2 from snapshot values in order
    snapshot_values = [v for k, v in snapshot.items() if k !=
                       "kind" and v is not None]
    for i, val in enumerate(snapshot_values, start=1):
        ctx[f"field_{i}"] = str(val)
    return ctx


def _render_template_body(body: str, context: dict) -> str:
    """Render template_body with Jinja2. Missing keys become empty string."""
    if not body.strip():
        return ""
    t = Template(body)
    return t.render(**{k: (v if v is not None else "") for k, v in context.items()})


def _fill_docx_paragraphs(doc: docx.Document, context: dict) -> None:
    """Replace {{ key }} placeholders in all paragraphs of a docx Document."""
    t = Template("")
    for para in doc.paragraphs:
        if not para.text:
            continue
        try:
            rendered = Template(para.text).render(
                **{k: (v if v is not None else "") for k, v in context.items()})
            if rendered != para.text:
                para.clear()
                para.add_run(rendered)
        except Exception:
            pass  # leave paragraph unchanged on render error


def _text_to_docx_bytes(text: str) -> bytes:
    """Create a minimal .docx in memory containing the given text."""
    doc = docx.Document()
    for line in (text or "").splitlines():
        doc.add_paragraph(line)
    if not (text or "").strip():
        doc.add_paragraph("")
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


def _get_system_user():
    """Return a user for background-created documents (e.g. first superuser)."""
    return User.objects.filter(is_superuser=True).first()


def _generate_document_for_run(run: ProcessRun, template: ProcessTemplate, snapshot: dict) -> None:
    doc_template = template.document_template
    if not doc_template:
        return

    system_user = _get_system_user()
    if not system_user:
        logger.warning(
            "Cannot generate document for run id=%s: no system user (superuser) found",
            run.id,
        )
        return

    category = doc_template.category
    if not category:
        category = DocumentCategory.objects.first()
    if not category:
        logger.warning(
            "Cannot generate document for run id=%s: no document category (template has no category and none in DB)",
            run.id,
        )
        return

    context = _build_document_context(run, snapshot)
    content_bytes: bytes | None = None
    ext = ".docx"
    title_suffix = f"Run #{run.id}"

    if doc_template.template_file:
        name = getattr(doc_template.template_file, "name", "") or ""
        if (name and name.lower().endswith(".docx")):
            try:
                with doc_template.template_file.open("rb") as fh:
                    doc = docx.Document(io.BytesIO(fh.read()))
                _fill_docx_paragraphs(doc, context)
                buf = io.BytesIO()
                doc.save(buf)
                buf.seek(0)
                content_bytes = buf.read()
            except Exception as e:
                logger.warning(
                    "Failed to fill docx template for run id=%s: %s",
                    run.id,
                    e,
                )
                if doc_template.template_body:
                    rendered = _render_template_body(
                        doc_template.template_body, context)
                    content_bytes = _text_to_docx_bytes(rendered)
        elif doc_template.template_body:
            rendered = _render_template_body(
                doc_template.template_body, context)
            content_bytes = _text_to_docx_bytes(rendered)
    elif doc_template.template_body:
        rendered = _render_template_body(doc_template.template_body, context)
        content_bytes = _text_to_docx_bytes(rendered)

    if not content_bytes:
        logger.warning(
            "No content for document generation run id=%s template id=%s (missing template_body and valid template_file)",
            run.id,
            template.id,
        )
        return

    title = f"{doc_template.name} – {title_suffix}"
    doc_file = DocumentFile(
        category=category,
        title=title,
        uploaded_by=system_user,
        valid_from=run.scheduled_for,
        valid_until=run.valid_until,
    )
    doc_file.file.save(
        f"process_run_{run.id}_{doc_template.id}{ext}",
        ContentFile(content_bytes),
        save=True,
    )
    doc_file.save()

    ProcessRunDocument.objects.create(
        process_run=run,
        document_file=doc_file,
        usage_kind=ProcessRunDocument.USAGE_REPORT,
    )
    logger.info(
        "Generated document id=%s for run id=%s from template id=%s",
        doc_file.id,
        run.id,
        template.id,
    )


def _send_email_for_template(
    template: ProcessTemplate,
    binding: ProcessBinding,
    snapshot: dict,
    run: ProcessRun | None = None,
) -> None:
    recipient = _resolve_email_recipient(template, binding)
    if not recipient:
        logger.warning(
            "No email recipient for ProcessTemplate id=%s", template.id)
        return

    subject = template.email_subject_template or "Process notification"
    body = template.email_body_template or ""
    if run is not None:
        context = _build_document_context(run, snapshot)
        subject = _render_template_body(subject, context)
        body = _render_template_body(body, context)
    try:
        sent = get_email_sender().send(
            recipients=[recipient],
            subject=subject,
            body=body,
            fail_silently=True,
        )
        if not sent:
            logger.warning(
                "Email send returned False for ProcessTemplate id=%s", template.id
            )
    except Exception as e:
        logger.exception(
            "Failed to send email for ProcessTemplate id=%s: %s", template.id, e
        )


def _resolve_email_recipient(template: ProcessTemplate, binding: ProcessBinding) -> str | None:
    if template.email_to_kind == ProcessTemplate.EMAIL_TO_CUSTOM:
        return template.custom_email_recipient or None
    if template.email_to_kind == ProcessTemplate.EMAIL_TO_CLIENT_MAIN and binding.client_company_id:
        return binding.client_company.email or None
    if template.email_to_kind == ProcessTemplate.EMAIL_TO_EMPLOYEE and binding.employee_id:
        return binding.employee.email or None
    return None


def run_on_completed_trigger(run: ProcessRun) -> None:
    templates = ProcessTemplate.objects.filter(
        process_type_id=run.process_type_id,
        trigger=ProcessTemplate.TRIGGER_ON_COMPLETED,
    )
    for template in templates:
        logger.info(
            "ON_COMPLETED trigger: run_id=%s process_type=%s template_id=%s template=%s",
            run.id,
            run.process_type_id,
            template.id,
            template,
        )


def run_on_expired_trigger(run: ProcessRun) -> None:
    binding = run.process_binding
    snapshot = run.subject_snapshot or {}
    templates = ProcessTemplate.objects.filter(
        process_type_id=run.process_type_id,
        trigger=ProcessTemplate.TRIGGER_ON_EXPIRED,
    )
    for template in templates:
        if template.send_email:
            _send_email_for_template(template, binding, snapshot, run=run)


def run_expired_reminders() -> None:
    today = date.today()
    runs = (
        ProcessRun.objects.filter(
            status=ProcessRun.STATUS_COMPLETED,
            valid_until__lt=today,
        )
        .select_related("process_binding", "process_type")
    )
    n = 0
    for run in runs:
        try:
            run_on_expired_trigger(run)
            n += 1
        except Exception as e:
            logger.exception(
                "Failed to run ON_EXPIRED for run id=%s: %s", run.id, e
            )
    if n:
        logger.info("run_expired_reminders: processed %s expired run(s)", n)
